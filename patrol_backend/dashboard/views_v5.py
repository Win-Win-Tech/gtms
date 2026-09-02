"""v5 dashboard APIs — site-wise scope. Live dashboard URLs unchanged."""

from django.http import HttpResponse
from django.utils import timezone
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action, parser_classes
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from authapp.site_access import assert_caller_can_access_site, caller_can_access_site, get_site_or_error
from patrol_backend.utils.timezone_utils import get_user_today, get_user_timezone_from_request
from scheduler.daily_site import apply_select_site, site_ids_payload
from scheduler.models import Assignment

from .site_filter import (
    assert_attendance_checkin_in_scope,
    assert_site_in_scope,
    assignment_attendance_site_q,
    inject_validated_site_query_params,
    resolve_dashboard_location_id,
    resolve_dashboard_site_scope,
    site_scope_internal_kwargs,
)
from .views import (
    AttendanceCheckinViewSet,
    AttendanceCheckinV4ExportView,
    AttendanceCheckinV4ListView,
    AttendanceCheckinV4PdfExportView,
    DashboardCheckInReportExcelViewV2,
    DashboardCheckInReportPdfViewV2,
    DashboardCheckInReportViewV2,
    MonthlyAttendanceExcelViewSetV2,
    MonthlyAttendanceSummaryViewSetV2,
    _attendance_v3_refresh_saved_fields,
    _attendance_v3_shift_window_utc,
    _get_checkin_report_data_v2,
    generate_attendance_v4_excel_report_internal,
)
from dashboard.models import AttendanceCheckin, CheckInLog
from dashboard.serializers import AttendanceCheckinSerializer
from patrol_backend.utils.attendance_v5 import get_open_checkin_log, resolve_v5_punch_site
from dashboard.attendance_detail_pdf import generate_attendance_v4_pdf_report_internal
from django.core.files.base import ContentFile


def _v5_error(exc):
    if isinstance(exc, ValidationError):
        detail = exc.detail
        if isinstance(detail, dict):
            return Response(detail, status=status.HTTP_400_BAD_REQUEST)
        return Response({"error": detail}, status=status.HTTP_400_BAD_REQUEST)
    if isinstance(exc, PermissionDenied):
        return Response({"error": str(exc)}, status=status.HTTP_403_FORBIDDEN)
    raise exc


class ShiftTodayV5(APIView):
    """
    Copy of GET /dashboard/attendance/shift_today_v3/ plus posted site fields.
    Mobile: use instead of shift_today_v3 when site-wise is enabled.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from dashboard.views import AttendanceCheckinViewSet

        view = AttendanceCheckinViewSet()
        view.request = request
        view.format_kwarg = None
        response = view.shift_today_v3(request)
        if response.status_code != status.HTTP_200_OK:
            return response

        user_tz = get_user_timezone_from_request(request)
        on_date = get_user_today(user_tz)
        data = dict(response.data)
        data.update(site_ids_payload(request.user, on_date))
        return Response(data, status=status.HTTP_200_OK)


class CreateAssignmentV5(APIView):
    """
    Copy of POST /dashboard/attendance/create_assignment/ plus optional site_id.
    Mobile self-assign: guard picks shift (+ template) and sends site_id.
    Site write uses attendance rules (daily vs cache), same as POST /auth/v5/my-sites/.
    """

    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        site_id = request.data.get("site_id") or None
        guard_id = request.data.get("guard_id")
        site = None
        try:
            if site_id:
                site = get_site_or_error(site_id)
                assert_caller_can_access_site(request.user, site)
                if guard_id:
                    from authapp.models import User

                    guard = User.objects.filter(id=guard_id).first()
                    if guard and not caller_can_access_site(guard, site) and not guard.is_superuser:
                        raise ValidationError({"site_id": "This guard is not assigned to that site."})

            from dashboard.views import AttendanceCheckinViewSet

            view = AttendanceCheckinViewSet()
            view.request = request
            view.format_kwarg = None
            response = view.create_assignment(request)
            if response.status_code == 201 and site_id and response.data.get("assignment_id"):
                assignment = Assignment.objects.filter(id=response.data["assignment_id"]).select_related("guard").first()
                if assignment:
                    user_tz = get_user_timezone_from_request(request)
                    on_date = assignment.start_date or get_user_today(user_tz)
                    selected = apply_select_site(
                        assignment.guard,
                        site_id,
                        on_date,
                        assignment=assignment,
                        caller=request.user,
                    )
                    data = dict(response.data)
                    data["site_id"] = str(site_id)
                    data["site_name"] = site.name
                    data["assigned_site_id"] = str(site_id)
                    data["assigned_site_name"] = site.name
                    if selected:
                        data["last_selected_site_id"] = selected.get("id")
                        data["last_selected_site_name"] = selected.get("name")
                    return Response(data, status=status.HTTP_201_CREATED)
            return response
        except ValidationError as e:
            return Response({"error": e.detail}, status=status.HTTP_400_BAD_REQUEST)
        except PermissionDenied as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)


class AttendanceCheckinV4ListViewV5(AttendanceCheckinV4ListView):
    """GET /dashboard/api/attendance_v5/ — v4 list with enforced site access."""

    def get_queryset(self):
        extra_q = inject_validated_site_query_params(self.request)
        qs = super().get_queryset()
        if extra_q is not None:
            qs = qs.filter(extra_q)
        return qs

    def list(self, request, *args, **kwargs):
        try:
            return super().list(request, *args, **kwargs)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)


class AttendanceCheckinV4ExportViewV5(AttendanceCheckinV4ExportView):
    """GET /dashboard/api/attendance/export_v5/ — Excel with enforced site access."""

    def get(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        params = request.query_params
        try:
            result = generate_attendance_v4_excel_report_internal(
                date_filter=params.get("date_filter", "today"),
                start_date=params.get("start_date"),
                end_date=params.get("end_date"),
                guard_id=params.get("guard"),
                location_id=params.get("location"),
                shift_id=params.get("shift"),
                status_filter=params.get("status"),
                defaulters=params.get("defaulters") == "true",
                request=request,
                search=params.get("search"),
                role=params.get("role"),
                site_id=scope.get("site_id"),
                site_scope_extra_q=scope.get("site_scope_extra_q"),
            )
            with open(result["file_path"], "rb") as f:
                content = f.read()
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            if params.get("date_filter") == "custom":
                fname = f"attendance_report_v5_{params.get('start_date', '')}_{params.get('end_date', '')}.xlsx"
            else:
                fname = f"attendance_report_v5_{timezone.now().strftime('%Y%m%d')}.xlsx"
            response["Content-Disposition"] = f'attachment; filename="{fname}"'
            return response
        except ValueError as e:
            return Response({"error": str(e)}, status=400)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)
        except Exception as e:
            from dashboard.views import logger

            logger.error(f"[ATTENDANCE_EXPORT_V5_API] Exception: {str(e)}", exc_info=True)
            return Response({"error": f"Failed to generate v5 report: {str(e)}"}, status=500)


class AttendanceCheckinV4PdfExportViewV5(AttendanceCheckinV4PdfExportView):
    """GET /dashboard/api/attendance/export_v5_pdf/ — PDF with enforced site access."""

    def get(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        params = request.query_params
        try:
            result = generate_attendance_v4_pdf_report_internal(
                date_filter=params.get("date_filter", "today"),
                start_date=params.get("start_date"),
                end_date=params.get("end_date"),
                guard_id=params.get("guard"),
                location_id=params.get("location"),
                shift_id=params.get("shift"),
                status_filter=params.get("status"),
                defaulters=params.get("defaulters") == "true",
                request=request,
                search=params.get("search"),
                role=params.get("role"),
                site_id=scope.get("site_id"),
                site_scope_extra_q=scope.get("site_scope_extra_q"),
            )
            content = result.get("pdf_bytes")
            if not content:
                with open(result["file_path"], "rb") as f:
                    content = f.read()
            response = HttpResponse(content, content_type="application/pdf")
            if params.get("date_filter") == "custom" and params.get("start_date"):
                fname = f"attendance_detail_v5_{params.get('start_date')}_{params.get('end_date', '')}.pdf"
            else:
                fname = f"attendance_detail_v5_{timezone.now().strftime('%Y%m%d')}.pdf"
            response["Content-Disposition"] = f'attachment; filename="{fname}"'
            return response
        except ValueError as e:
            return Response({"error": str(e)}, status=400)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)
        except Exception as e:
            from dashboard.views import logger

            logger.error(f"[ATTENDANCE_EXPORT_V5_PDF] Exception: {str(e)}", exc_info=True)
            return Response({"error": f"Failed to generate PDF report: {str(e)}"}, status=500)


class DashboardCheckInReportViewV5(DashboardCheckInReportViewV2):
    """GET /dashboard/dashboard-checkin-report/v5/ — check-in report with site access."""

    def get(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
            assignment_site_q = assignment_attendance_site_q(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        filter_type = request.query_params.get("filter", "today")
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")
        user_id = request.query_params.get("user_id")
        location_id = request.query_params.get("location_id")
        shift_id = request.query_params.get("shift_id")
        search = (request.query_params.get("search") or "").strip()
        role = (request.query_params.get("role") or "").strip()

        try:
            status_filter = (request.query_params.get("status") or "all").strip()
            report_data = _get_checkin_report_data_v2(
                filter_type=filter_type,
                start_date_str=start_date,
                end_date_str=end_date,
                user_id=user_id,
                location_id=location_id,
                shift_id=shift_id,
                request=request,
                search=search or None,
                role=role or None,
                site_id=scope.get("site_id"),
                assignment_site_q=assignment_site_q,
            )

            if status_filter and status_filter != "all":
                report_data = [r for r in report_data if r.get("status") == status_filter]

            for item in report_data:
                if item["expected_time"]:
                    if hasattr(item["expected_time"], "isoformat"):
                        item["expected_time"] = item["expected_time"].isoformat()
                    else:
                        item["expected_time"] = item["expected_time"].strftime("%Y-%m-%d %H:%M:%S")
                if item["actual_checkin_time"]:
                    if hasattr(item["actual_checkin_time"], "isoformat"):
                        item["actual_checkin_time"] = item["actual_checkin_time"].isoformat()
                    else:
                        item["actual_checkin_time"] = item["actual_checkin_time"].strftime(
                            "%Y-%m-%d %H:%M:%S"
                        )

            from dashboard.serializers import CheckInReportSerializer

            serializer = CheckInReportSerializer(report_data, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)

        except ValueError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response(
                {"error": f"An error occurred while generating the v5 report: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class DashboardCheckInReportExcelViewV5(DashboardCheckInReportExcelViewV2):
    """GET /dashboard/dashboard-checkin-report-excel/v5/ — Excel with site access."""

    def get(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
            assignment_site_q = assignment_attendance_site_q(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        filter_type = request.query_params.get("filter", "today")
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")
        user_id = request.query_params.get("user_id")
        location_id = request.query_params.get("location_id")
        shift_id = request.query_params.get("shift_id")
        search = (request.query_params.get("search") or "").strip()
        role = (request.query_params.get("role") or "").strip()
        status_filter = (request.query_params.get("status") or "").strip()

        try:
            from io import BytesIO

            from django.http import HttpResponse
            from django.utils import timezone
            from openpyxl import Workbook

            report_data = _get_checkin_report_data_v2(
                filter_type=filter_type,
                start_date_str=start_date,
                end_date_str=end_date,
                user_id=user_id,
                location_id=location_id,
                shift_id=shift_id,
                request=request,
                search=search or None,
                role=role or None,
                site_id=scope.get("site_id"),
                assignment_site_q=assignment_site_q,
            )

            if status_filter and status_filter != "all":
                report_data = [r for r in report_data if r.get("status") == status_filter]

            wb = Workbook()
            ws = wb.active
            ws.title = "QR Scan Patrol Report"

            def format_checklist_cell(report_item):
                answers = report_item.get("checklist_answers")
                if isinstance(answers, list) and len(answers) > 0:
                    lines = []
                    for a in answers:
                        label = (a or {}).get("label") or "Item"
                        checked = (a or {}).get("checked") is True
                        lines.append(f"{label} : {'✓' if checked else '✗'}")
                    return "\n".join(lines)
                return report_item.get("checklist_template_name") or ""

            headers = [
                "Date",
                "Name",
                "Emp Code",
                "Designation",
                "Shift Name",
                "Checkpoint Name",
                "Site",
                "Expected Time",
                "Actual Check-In Time",
                "Status",
                "Delay (minutes)",
                "Has Checklist",
                "Checklist",
                "Checklist Remarks",
            ]
            ws.append(headers)

            for item in report_data:
                ws.append([
                    item["date"],
                    item["guard_name"],
                    item.get("employee_code") or "",
                    item.get("designation") or "",
                    item["shift_name"],
                    item["checkpoint_name"],
                    item.get("site_name") or "",
                    item["expected_time"].strftime("%Y-%m-%d %H:%M") if item["expected_time"] else "",
                    item["actual_checkin_time"].strftime("%Y-%m-%d %H:%M")
                    if item["actual_checkin_time"]
                    else "",
                    item["status"],
                    item["delay_minutes"] if item["delay_minutes"] is not None else "",
                    "Yes" if item.get("has_checklist") else "No",
                    format_checklist_cell(item),
                    item.get("checklist_remarks") or "",
                ])

            try:
                from openpyxl.styles import Alignment

                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=12, max_col=12):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.column_dimensions["L"].width = 55
                ws.column_dimensions["M"].width = 35
            except Exception:
                pass

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            excel_content = buffer.getvalue()

            response = HttpResponse(
                excel_content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            if filter_type == "custom" and start_date and end_date:
                filename = f"checkin_report_v5_{start_date}_{end_date}.xlsx"
            else:
                filename = f"checkin_report_v5_{timezone.now().strftime('%Y%m%d')}.xlsx"
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            return response
        except ValueError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response(
                {"error": f"An error occurred while generating the v5 Excel report: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class DashboardCheckInReportPdfViewV5(DashboardCheckInReportPdfViewV2):
    """GET /dashboard/dashboard-checkin-report-pdf/v5/ — PDF with site access."""

    def get(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
            assignment_site_q = assignment_attendance_site_q(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        filter_type = request.query_params.get("filter", "today")
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")
        user_id = request.query_params.get("user_id")
        location_id = request.query_params.get("location_id")
        shift_id = request.query_params.get("shift_id")
        search = (request.query_params.get("search") or "").strip()
        role = (request.query_params.get("role") or "").strip()
        status_filter = (request.query_params.get("status") or "").strip()

        try:
            report_data = _get_checkin_report_data_v2(
                filter_type=filter_type,
                start_date_str=start_date,
                end_date_str=end_date,
                user_id=user_id,
                location_id=location_id,
                shift_id=shift_id,
                request=request,
                search=search or None,
                role=role or None,
                site_id=scope.get("site_id"),
                assignment_site_q=assignment_site_q,
            )
            if status_filter and status_filter != "all":
                report_data = [r for r in report_data if r.get("status") == status_filter]

            from dashboard.checkin_report_pdf import generate_checkin_report_pdf

            content, filename = generate_checkin_report_pdf(
                report_data,
                request,
                filter_type,
                start_date,
                end_date,
                location_id,
            )
            if not filename.startswith("checkin_report_v5"):
                filename = filename.replace("checkin_report", "checkin_report_v5", 1)
            response = HttpResponse(content, content_type="application/pdf")
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            return response
        except ImportError as e:
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except ValueError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response(
                {"error": f"An error occurred while generating the check-in PDF report: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class MonthlyAttendanceSummaryViewSetV5(MonthlyAttendanceSummaryViewSetV2):
    """GET /dashboard/attendance-summary-v5/summary/ — monthly grid with site access."""

    @action(detail=False, methods=["get"])
    def summary(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        p = request.query_params
        try:
            from dashboard.views import _get_monthly_attendance_summary_data_v2, logger

            res = _get_monthly_attendance_summary_data_v2(
                month=p.get("month"),
                start_date_str=p.get("start_date"),
                end_date_str=p.get("end_date"),
                location_id=p.get("location_id"),
                site_id=scope.get("site_id"),
                user_id=p.get("user_id"),
                search=p.get("search"),
                role=p.get("role"),
                request=request,
                site_scope_extra_q=scope.get("site_scope_extra_q"),
            )
            return Response(res["summary_data"])
        except ValueError as e:
            return Response({"error": str(e)}, status=400)
        except Exception as e:
            from dashboard.views import logger

            logger.error(f"[MONTHLY_ATTENDANCE_SUMMARY_V5_API] {str(e)}", exc_info=True)
            return Response({"error": str(e)}, status=500)


class MonthlyAttendanceExcelViewSetV5(MonthlyAttendanceExcelViewSetV2):
    """GET /dashboard/attendance-excel-v5/export_excel/ — monthly Excel with site access."""

    @action(detail=False, methods=["get"])
    def export_excel(self, request):
        try:
            scope = site_scope_internal_kwargs(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

        p = request.query_params
        try:
            from dashboard.views import generate_monthly_attendance_summary_excel_internal_v2

            res = generate_monthly_attendance_summary_excel_internal_v2(
                month=p.get("month"),
                start_date_str=p.get("start_date"),
                end_date_str=p.get("end_date"),
                request=request,
                location_id=p.get("location_id"),
                site_id=scope.get("site_id"),
                user_id=p.get("user_id"),
                search=p.get("search"),
                role=p.get("role"),
                include_location_column=request.user.is_superuser,
                site_scope_extra_q=scope.get("site_scope_extra_q"),
            )
            with open(res["file_path"], "rb") as f:
                content = f.read()
            response = HttpResponse(
                content,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = (
                f'attachment; filename="attendance_summary_v5_{res["start_date"].strftime("%Y%m%d")}.xlsx"'
            )
            return response
        except ValueError as e:
            return Response({"error": str(e)}, status=400)
        except Exception as e:
            return Response({"error": str(e)}, status=500)


def _filter_bulk_candidate_payload(data, request):
    """Trim bulk-entry candidate lists to header site user scope."""
    effective_site_id, extra_q = resolve_dashboard_site_scope(request)
    if not effective_site_id and not extra_q:
        return data

    from authapp.models import User, UserSite
    from authapp.site_access import allowed_site_ids

    if effective_site_id:
        allowed_users = set(
            str(x)
            for x in UserSite.objects.filter(site_id=effective_site_id).values_list("user_id", flat=True)
        )
    else:
        allowed = allowed_site_ids(request.user)
        allowed_users = set(
            str(x)
            for x in UserSite.objects.filter(site_id__in=allowed).values_list("user_id", flat=True)
        )

    location_id = resolve_dashboard_location_id(request)
    if location_id:
        org_wide = User.objects.filter(all_org_sites=True, location_id=location_id, is_deleted=False)
        allowed_users |= {str(x) for x in org_wide.values_list("id", flat=True)}

    for key in ("no_shift", "shift_no_checkin", "weekoff"):
        if key in data and isinstance(data[key], list):
            data[key] = [row for row in data[key] if str(row.get("user_id")) in allowed_users]

    summary = data.get("summary") or {}
    summary["no_shift_count"] = len(data.get("no_shift") or [])
    summary["shift_no_checkin_count"] = len(data.get("shift_no_checkin") or [])
    summary["weekoff_count"] = len(data.get("weekoff") or [])
    summary["total_users"] = len(
        {row.get("user_id") for key in ("no_shift", "shift_no_checkin", "weekoff") for row in (data.get(key) or [])}
    )
    data["summary"] = summary
    return data


class AttendanceCheckinViewSetV5(viewsets.GenericViewSet):
    """Web attendance mutations with enforced site access. Does not re-expose live actions."""

    permission_classes = [permissions.IsAuthenticated]
    queryset = AttendanceCheckinViewSet.queryset

    def _live(self):
        view = AttendanceCheckinViewSet()
        view.request = self.request
        view.format_kwarg = self.format_kwarg
        return view

    def _parse_punch_coords(self, request):
        try:
            lat = float(request.data.get("latitude"))
            lon = float(request.data.get("longitude"))
        except (TypeError, ValueError):
            raise ValidationError({"error": "latitude and longitude are required"})
        return lat, lon

    def _v5_punch_response(self, request, attendance, matched_site, resolution, action_mode, status_code):
        data = AttendanceCheckinSerializer(attendance, context={"request": request}).data
        if isinstance(data, dict):
            org_location = attendance.org_location
            data["face_attendance"] = bool(getattr(org_location, "is_face_attendance_enabled", False))
            data["face_verified"] = bool(getattr(org_location, "is_face_attendance_enabled", False))
            data["site_id"] = str(matched_site.id) if matched_site else None
            data["site_name"] = matched_site.name if matched_site else None
            data["site_resolution"] = resolution
            data["mode"] = action_mode
        return Response(data, status=status_code)

    @action(detail=False, methods=["post"], url_path="checkin_v5")
    @parser_classes([MultiPartParser, FormParser])
    def checkin_v5(self, request):
        """
        Mobile check-in with explicit site_id (v5).
        site_id empty → nearest site within attendance_distance.
        """
        from patrol_backend.utils.attendance_resolve import (
            CheckinTooSoonAfterCheckout,
            apply_v4_attendance_after_log,
            build_log_window_filter,
            enforce_checkin_allowed_after_checkout,
            get_or_create_attendance_for_shift_day,
        )
        from patrol_backend.utils.face_utils import is_face_attendance_available, verify_user_face

        try:
            user = request.user
            user_tz = get_user_timezone_from_request(request)
            live = self._live()
            assignment, shift_start_date = live.get_today_assignment_v2(user, request)

            if not assignment:
                return Response({"message": "No shifts today"}, status=status.HTTP_400_BAD_REQUEST)

            shift = assignment.shift
            org_location = assignment.location
            lat, lon = self._parse_punch_coords(request)

            matched_site, _dist, resolution = resolve_v5_punch_site(
                site_id=request.data.get("site_id"),
                latitude=lat,
                longitude=lon,
                org_location_id=org_location.id,
                user=user,
            )

            search_start_utc, search_end_utc, _, _, _ = _attendance_v3_shift_window_utc(
                shift_start_date, shift, user_tz, location_id=getattr(org_location, "id", None)
            )
            log_filter = build_log_window_filter(
                user, assignment, shift, org_location, search_start_utc, search_end_utc
            )

            try:
                enforce_checkin_allowed_after_checkout(
                    user, assignment, shift, org_location, search_start_utc, search_end_utc
                )
            except CheckinTooSoonAfterCheckout as exc:
                return Response(
                    {
                        "error": f"Check-in allowed {exc.min_minutes} minute(s) after checkout",
                        "code": "checkin_too_soon_after_checkout",
                        "min_checkin_after_checkout_minutes": exc.min_minutes,
                        "remaining_seconds": exc.remaining,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            image_file = request.FILES.get("image") or request.FILES.get("checkin_image")
            raw_bytes = image_file.read() if image_file else None

            if getattr(org_location, "is_face_attendance_enabled", False):
                if not is_face_attendance_available():
                    return Response(
                        {
                            "error": "Face attendance is enabled for this location but face_recognition is not installed on the server.",
                            "hint": "See docs/FACE_ATTENDANCE_INSTALL.md",
                        },
                        status=status.HTTP_503_SERVICE_UNAVAILABLE,
                    )
                if not raw_bytes:
                    return Response({"error": "Face attendance requires an image"}, status=status.HTTP_400_BAD_REQUEST)
                ok, msg, _dist_face = verify_user_face(user, raw_bytes)
                if not ok:
                    return Response({"error": msg}, status=status.HTTP_400_BAD_REQUEST)

            img_name = "checkin.jpg"
            if image_file:
                img_name = getattr(image_file, "name", img_name) or img_name

            log = CheckInLog.objects.create(
                guard=user,
                assignment=assignment,
                shift=shift,
                org_location=org_location,
                type="checkin",
                latitude=lat,
                longitude=lon,
                site=matched_site,
            )
            if raw_bytes is not None:
                log.image.save(img_name, ContentFile(raw_bytes), save=True)

            attendance = get_or_create_attendance_for_shift_day(
                user, assignment, shift, org_location, shift_start_date
            )
            apply_v4_attendance_after_log(
                attendance=attendance,
                user=user,
                assignment=assignment,
                shift=shift,
                org_location=org_location,
                shift_start_date=shift_start_date,
                matched_site=matched_site,
                log_filter=log_filter,
                search_start_utc=search_start_utc,
                search_end_utc=search_end_utc,
                action_mode="checkin",
                raw_bytes=raw_bytes,
                img_name=img_name,
                user_tz=user_tz,
                refresh_fn=_attendance_v3_refresh_saved_fields,
                skip_sibling_reconcile=False,
            )

            return self._v5_punch_response(
                request, attendance, matched_site, resolution, "checkin", status.HTTP_201_CREATED
            )
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="checkout_v5")
    @parser_classes([MultiPartParser, FormParser])
    def checkout_v5(self, request):
        """
        Mobile check-out with explicit site_id (v5).
        site_id empty → nearest site within attendance_distance.
        Checkout site must match the open check-in site.
        """
        from patrol_backend.utils.attendance_resolve import (
            apply_v4_attendance_after_log,
            build_log_window_filter,
            get_or_create_attendance_for_shift_day,
        )
        from patrol_backend.utils.face_utils import is_face_attendance_available, verify_user_face

        try:
            user = request.user
            user_tz = get_user_timezone_from_request(request)
            live = self._live()
            assignment, shift_start_date = live.get_today_assignment_v2(user, request)

            if not assignment:
                return Response({"message": "No shifts today"}, status=status.HTTP_400_BAD_REQUEST)

            shift = assignment.shift
            org_location = assignment.location
            lat, lon = self._parse_punch_coords(request)

            search_start_utc, search_end_utc, _, _, _ = _attendance_v3_shift_window_utc(
                shift_start_date, shift, user_tz, location_id=getattr(org_location, "id", None)
            )
            log_filter = build_log_window_filter(
                user, assignment, shift, org_location, search_start_utc, search_end_utc
            )

            open_checkin = get_open_checkin_log(
                user, assignment, shift, org_location, search_start_utc, search_end_utc
            )
            if not open_checkin:
                attendance_probe = AttendanceCheckin.objects.filter(
                    guard=user,
                    assignment=assignment,
                    shift=shift,
                    org_location=org_location,
                    shift_date=shift_start_date,
                ).first()
                has_checkin = CheckInLog.objects.filter(**log_filter, type="checkin").exists()
                if not has_checkin and not (attendance_probe and attendance_probe.checkin_time):
                    return Response({"message": "Cannot checkout before checkin"}, status=status.HTTP_400_BAD_REQUEST)
                return Response(
                    {
                        "error": "No open check-in session found for checkout.",
                        "code": "no_open_checkin",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            checkin_site_id = open_checkin.site_id
            if not checkin_site_id:
                return Response(
                    {
                        "error": "Check-in site is missing. Please check in again using checkin_v5.",
                        "code": "checkin_site_missing",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            matched_site, _dist, resolution = resolve_v5_punch_site(
                site_id=request.data.get("site_id"),
                latitude=lat,
                longitude=lon,
                org_location_id=org_location.id,
                user=user,
            )

            if str(matched_site.id) != str(checkin_site_id):
                checkin_site_name = getattr(open_checkin.site, "name", None) or str(checkin_site_id)
                return Response(
                    {
                        "error": "Checkout site must match your check-in site.",
                        "code": "checkout_site_mismatch",
                        "checkin_site_id": str(checkin_site_id),
                        "checkin_site_name": checkin_site_name,
                        "checkout_site_id": str(matched_site.id),
                        "checkout_site_name": matched_site.name,
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            attendance = get_or_create_attendance_for_shift_day(
                user, assignment, shift, org_location, shift_start_date
            )

            image_file = request.FILES.get("image") or request.FILES.get("checkout_image")
            raw_bytes = image_file.read() if image_file else None

            if getattr(org_location, "is_face_attendance_enabled", False):
                if not is_face_attendance_available():
                    return Response(
                        {
                            "error": "Face attendance is enabled for this location but face_recognition is not installed on the server.",
                            "hint": "See docs/FACE_ATTENDANCE_INSTALL.md",
                        },
                        status=status.HTTP_503_SERVICE_UNAVAILABLE,
                    )
                if not raw_bytes:
                    return Response({"error": "Face attendance requires an image"}, status=status.HTTP_400_BAD_REQUEST)
                ok, msg, _dist_face = verify_user_face(user, raw_bytes)
                if not ok:
                    return Response({"error": msg}, status=status.HTTP_400_BAD_REQUEST)

            img_name = "checkout.jpg"
            if image_file:
                img_name = getattr(image_file, "name", img_name) or img_name

            log = CheckInLog.objects.create(
                guard=user,
                assignment=assignment,
                shift=shift,
                org_location=org_location,
                type="checkout",
                latitude=lat,
                longitude=lon,
                site=matched_site,
            )
            if raw_bytes is not None:
                log.image.save(img_name, ContentFile(raw_bytes), save=True)

            apply_v4_attendance_after_log(
                attendance=attendance,
                user=user,
                assignment=assignment,
                shift=shift,
                org_location=org_location,
                shift_start_date=shift_start_date,
                matched_site=matched_site,
                log_filter=log_filter,
                search_start_utc=search_start_utc,
                search_end_utc=search_end_utc,
                action_mode="checkout",
                raw_bytes=raw_bytes,
                img_name=img_name,
                user_tz=user_tz,
                refresh_fn=_attendance_v3_refresh_saved_fields,
                skip_sibling_reconcile=False,
            )

            return self._v5_punch_response(
                request, attendance, matched_site, resolution, "checkout", status.HTTP_200_OK
            )
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["get"], url_path="v5/bulk-entry-candidates")
    def bulk_entry_candidates_v5(self, request):
        try:
            site_scope_internal_kwargs(request)
            response = self._live().bulk_entry_candidates_v2(request)
            if response.status_code != status.HTTP_200_OK:
                return response
            data = _filter_bulk_candidate_payload(dict(response.data), request)
            return Response(data, status=status.HTTP_200_OK)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="bulk-entry_v5")
    def bulk_entry_v5(self, request):
        try:
            assert_site_in_scope(request, request.data.get("site_id"))
            return self._live().bulk_entry(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="bulk-weekoff_v5")
    def bulk_weekoff_v5(self, request):
        try:
            assert_site_in_scope(request, request.data.get("site_id"))
            return self._live().bulk_weekoff(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="monthly-cell-action_v5")
    def monthly_cell_action_v5(self, request):
        try:
            assert_site_in_scope(request, request.data.get("site_id"))
            return self._live().monthly_cell_action(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="force-checkout_v5")
    @parser_classes([MultiPartParser, FormParser])
    def force_checkout_v5(self, request):
        try:
            attendance_id = request.data.get("attendance_id")
            if attendance_id:
                from dashboard.models import AttendanceCheckin

                attendance = AttendanceCheckin.objects.filter(id=attendance_id).first()
                assert_attendance_checkin_in_scope(request, attendance)
            assert_site_in_scope(request, request.data.get("site_id"))
            return self._live().force_checkout_v3(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="edit-boundary_v5")
    def edit_boundary_v5(self, request):
        try:
            attendance_id = request.data.get("attendance_id")
            if attendance_id:
                from dashboard.models import AttendanceCheckin

                attendance = AttendanceCheckin.objects.filter(id=attendance_id).first()
                assert_attendance_checkin_in_scope(request, attendance)
            return self._live().edit_boundary_v3(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="add-punch_v5")
    def add_punch_v5(self, request):
        try:
            assert_site_in_scope(request, request.data.get("site_id"))
            attendance_id = request.data.get("attendance_id")
            if attendance_id:
                from dashboard.models import AttendanceCheckin

                attendance = AttendanceCheckin.objects.filter(id=attendance_id).first()
                assert_attendance_checkin_in_scope(request, attendance)
            return self._live().add_punch_v3(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="edit-punch_v5")
    def edit_punch_v5(self, request):
        try:
            from dashboard.models import AttendanceCheckin, CheckInLog

            log_id = request.data.get("log_id")
            if log_id:
                log = CheckInLog.objects.select_related("site").filter(id=log_id).first()
                if log and log.site_id:
                    assert_site_in_scope(request, str(log.site_id))
                if log:
                    attendance = AttendanceCheckin.objects.filter(
                        guard_id=log.guard_id,
                        assignment_id=log.assignment_id,
                        shift_id=log.shift_id,
                        org_location_id=log.org_location_id,
                    ).order_by("-shift_date").first()
                    assert_attendance_checkin_in_scope(request, attendance)
            return self._live().edit_punch_v3(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)

    @action(detail=False, methods=["post"], url_path="delete-punch_v5")
    def delete_punch_v5(self, request):
        try:
            from dashboard.models import CheckInLog

            log_id = request.data.get("log_id")
            if log_id:
                log = CheckInLog.objects.select_related("site").filter(id=log_id).first()
                if log and log.site_id:
                    assert_site_in_scope(request, str(log.site_id))
                if log:
                    from dashboard.models import AttendanceCheckin

                    attendance = AttendanceCheckin.objects.filter(
                        guard_id=log.guard_id,
                        assignment_id=log.assignment_id,
                        shift_id=log.shift_id,
                        org_location_id=log.org_location_id,
                    ).order_by("-shift_date").first()
                    assert_attendance_checkin_in_scope(request, attendance)
            return self._live().delete_punch_v3(request)
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)
