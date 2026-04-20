import uuid

from rest_framework import viewsets, status as http_status
from django.db.models import Q, F
from .models import Location, Shift, Assignment, Checkpoint, SiteSetting
from checkin.models import CheckIn
from .serializers import (
    LocationSerializer,
    ShiftSerializer,
    AssignmentSerializer,
    CheckpointSerializer,
    SiteSettingSerializer
)
from rest_framework.decorators import action
from rest_framework.response import Response
from django.utils.timezone import now, make_aware
from rest_framework.exceptions import ValidationError
from datetime import datetime, timedelta
import pytz
import logging
from calendar import monthrange
from collections import defaultdict
from rest_framework.decorators import action
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from patrol_backend.utils.timezone_utils import (
    get_user_timezone_from_request,
    get_user_today,
    get_user_now,
    combine_date_time_in_user_tz,
    to_user_timezone,
    convert_date_range_to_utc
)

logger = logging.getLogger(__name__)

class LocationViewSet(viewsets.ModelViewSet):
    queryset = Location.objects.all()
    serializer_class = LocationSerializer
    
    def perform_create(self, serializer):
        # Save the new location
        instance = serializer.save(created_by=self.request.user)
        
        # Explicitly duplicate default roles for this new location
        try:
            from authapp.models import Role
            global_roles = Role.objects.filter(location__isnull=True)
            for g_role in global_roles:
                Role.objects.get_or_create(
                    name=g_role.name,
                    location=instance,
                    defaults={
                        'is_default': g_role.is_default,
                        'is_allow_webapp': g_role.is_allow_webapp,
                        'pages': g_role.pages
                    }
                )
            logger.info(f"Successfully duplicated default roles for new organization: {instance.name}")
        except Exception as e:
            logger.error(f"Failed to duplicate roles for location {instance.id}: {str(e)}")


    def get_queryset(self):
        try:
            queryset = Location.objects.filter(is_deleted=False)
            name = self.request.query_params.get('name')
            address = self.request.query_params.get('address')

            if name:
                queryset = queryset.filter(name__icontains=name)
            if address:
                queryset = queryset.filter(address__icontains=address)

            return queryset
        except Exception as e:
            logger.error(f"Error filtering locations: {e}", exc_info=True)
            return Location.objects.none()


class ShiftViewSet(viewsets.ModelViewSet):
    queryset = Shift.objects.all()
    serializer_class = ShiftSerializer

    @action(detail=False, methods=['get'], url_path='by-location/(?P<location_id>[^/.]+)')
    def by_location(self, request, location_id=None):
        try:
            shifts = Shift.objects.filter(location_id=location_id, is_deleted=False)
            serializer = self.get_serializer(shifts, many=True)
            return Response(serializer.data)
        except Exception as e:
            logger.error(f"Error fetching shifts by location: {e}", exc_info=True)
            return Response({'error': 'Failed to retrieve shifts.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class AssignmentViewSet(viewsets.ModelViewSet):
    queryset = Assignment.objects.all()
    serializer_class = AssignmentSerializer

    def clean(self):
        try:
            super().clean()
            if self.end_date < self.start_date:
                raise ValidationError("End date cannot be before start date.")

            overlapping = Assignment.objects.filter(
                guard=self.guard,
                start_date__lte=self.end_date,
                end_date__gte=self.start_date,
            ).exclude(id=self.id)

            if overlapping.exists():
                raise ValidationError(f"Guard {self.guard} already has an overlapping assignment.")
        except ValidationError as ve:
            raise ve
        except Exception as e:
            logger.error(f"Error during assignment validation: {e}", exc_info=True)
            raise ValidationError("Unexpected error during assignment validation.")

    @action(detail=False, methods=['post'], url_path='bulk-create')
    def bulk_create(self, request):
        assignments_data = request.data

        if not isinstance(assignments_data, list):
            return Response({"error": "Expected a list of assignments"}, status=status.HTTP_400_BAD_REQUEST)

        created = []
        errors = []

        for idx, data in enumerate(assignments_data):
            serializer = self.get_serializer(data=data)
            if serializer.is_valid():
                serializer.save()
                created.append(serializer.data)
            else:
                errors.append({
                    "index": idx,
                    "errors": serializer.errors,
                    "data": data
                })

        if errors:
            return Response({
                "created": created,
                "errors": errors
            }, status=status.HTTP_207_MULTI_STATUS)

        return Response({"created": created}, status=status.HTTP_201_CREATED)


    # @action(detail=False, methods=['get'], url_path='upcoming-checkpoints/(?P<user_id>[^/.]+)')
    # def upcoming_checkpoints(self, request, user_id=None):
    #     try:
    #         # 1. Timezone and Date Range Setup
    #         user_tz = get_user_timezone_from_request(request)
    #         today_initial = get_user_today(user_tz)
    #         yesterday_initial = today_initial - timedelta(days=1)

    #         from patrol_backend.utils.timezone_utils import convert_date_range_to_utc
    #         start_dt_user = datetime.combine(today_initial, datetime.min.time())
    #         end_dt_user = datetime.combine(today_initial, datetime.max.time())
    #         start_utc, end_utc = convert_date_range_to_utc(start_dt_user.date(), end_dt_user.date(), user_tz)
    #         end_utc = end_utc + timedelta(days=1)

    #         # 2. Fetch Assignments (Today + Overnight carry-overs)
    #         assignments = list(Assignment.objects.filter(
    #             start_date__lte=end_utc.date(),
    #             end_date__gte=start_utc.date(),
    #             guard_id=user_id
    #         ).select_related('shift', 'shift__location'))

    #         yesterday_assignments = list(Assignment.objects.filter(
    #             guard_id=user_id,
    #             end_date=yesterday_initial,
    #             shift__end_time__lte=F('shift__start_time')
    #         ).select_related('shift', 'shift__location'))

    #         all_assignments = assignments + yesterday_assignments

    #         # 3. COLLECT ALL SCANS (Fetch all scans in the period to match)
    #         all_checkins = list(CheckIn.objects.filter(
    #             guard_id=user_id,
    #             timestamp__gte=start_utc,
    #             timestamp__lt=end_utc
    #         ).order_by('timestamp'))

    #         used_checkin_ids = set()  # To track consumed physical scans
    #         result = []

    #         # 4. Process Each Assignment
    #         for assignment in all_assignments:
    #             location_id = assignment.shift.location_id if assignment.shift and assignment.shift.location else None
    #             assignment_tz = get_user_timezone_from_request(request, location_id=location_id) if location_id else user_tz

    #             today = get_user_today(assignment_tz)
    #             user_now = get_user_now(assignment_tz)
    #             user_now_utc = user_now.astimezone(pytz.UTC)
    #             yesterday = today - timedelta(days=1)

    #             shift = assignment.shift
    #             is_overnight = shift.end_time <= shift.start_time
    #             assignment_ended_yesterday = assignment.end_date == yesterday

    #             if assignment_ended_yesterday and is_overnight:
    #                 if user_now.time() >= shift.end_time:
    #                     continue

    #             # 5. Process Checkpoints in this Assignment
    #             for cp in assignment.checkpoints:
    #                 checkpoint_id_str = cp.get('checkpoint_id')
    #                 try:
    #                     checkpoint_id = uuid.UUID(checkpoint_id_str)
    #                 except (ValueError, TypeError):
    #                     continue

    #                 checkpoint_obj = Checkpoint.objects.filter(id=checkpoint_id).first()
    #                 if not checkpoint_obj:
    #                     continue

    #                 try:
    #                     checkpoint_time = datetime.strptime(cp['time'], '%H:%M').time()

    #                     # Determine actual calendar date for this instance
    #                     if is_overnight:
    #                         if checkpoint_time >= shift.start_time:
    #                             if assignment_ended_yesterday:
    #                                 continue
    #                             checkpoint_date = today
    #                         else:
    #                             if user_now.time() < shift.end_time:
    #                                 checkpoint_date = today
    #                             else:
    #                                 checkpoint_date = today + timedelta(days=1)
    #                                 if today > assignment.end_date:
    #                                     continue
    #                     else:
    #                         checkpoint_date = today

    #                     # Expected time in UTC
    #                     expected_dt_utc = combine_date_time_in_user_tz(checkpoint_date, checkpoint_time, assignment_tz)

    #                     # 6. SCAN CONSUMPTION MATCHING (Physical scans only)
    #                     is_checked_in = False
    #                     for scan in all_checkins:
    #                         # Only match physical scans (synced=True)
    #                         if (scan.checkpoint_id == checkpoint_id and
    #                             getattr(scan, 'synced', True) == True and
    #                             scan.id not in used_checkin_ids and
    #                             abs(scan.timestamp - expected_dt_utc) <= timedelta(minutes=30)):

    #                             is_checked_in = True
    #                             used_checkin_ids.add(scan.id)
    #                             break

    #                     # 7. OVERDUE & SYNC LOGIC
    #                     is_overdue = user_now_utc > (expected_dt_utc + timedelta(minutes=15))

    #                     if is_checked_in:
    #                         status = 'completed'
    #                         synced = True
    #                     elif is_overdue:
    #                         status = 'completed'  # Per requirement: status is 'completed' for missed
    #                         synced = False        # But synced is false

    #                         # Optional: Auto-create missed record in DB if it doesn't exist
    #                         CheckIn.objects.get_or_create(
    #                             guard_id=user_id,
    #                             shift_id=shift.id,
    #                             checkpoint_id=checkpoint_id,
    #                             timestamp=expected_dt_utc,
    #                             defaults={
    #                                 'latitude': checkpoint_obj.latitude,
    #                                 'longitude': checkpoint_obj.longitude,
    #                                 'synced': False
    #                             }
    #                         )
    #                     else:
    #                         status = 'pending'
    #                         synced = None

    #                     # Convert UTC time to local for mobile display
    #                     display_time = expected_dt_utc.astimezone(assignment_tz).strftime('%H:%M')

    #                     # NEW: internal sort key (fixes overnight ordering)
    #                     sort_ts = expected_dt_utc.timestamp()

    #                 except Exception as e:
    #                     logger.warning(f"[UPCOMING_CHECKPOINTS_API] Error: {e}")
    #                     status = 'pending'
    #                     synced = None
    #                     display_time = cp.get('time')
    #                     sort_ts = float('inf')  # push bad records to end

    #                 result.append({
    #                     'assign_id': assignment.id,
    #                     'checkpoint_id': checkpoint_id,
    #                     'label': checkpoint_obj.label,
    #                     'time': display_time,
    #                     'lat': checkpoint_obj.latitude,
    #                     'lon': checkpoint_obj.longitude,
    #                     'qr': checkpoint_obj.data,
    #                     'shift_id': shift.id,
    #                     'status': status,
    #                     'synced': synced,

    #                     # NEW: used only for sorting
    #                     'sort_ts': sort_ts
    #                 })

    #         # 8. FINAL SORTING BY TIME (FIXED)
    #         result.sort(key=lambda x: x['sort_ts'])

    #         # OPTIONAL: remove internal field
    #         for r in result:
    #             r.pop('sort_ts', None)

    #         return Response(result)

    #     except Exception as e:
    #         logger.error(f"[UPCOMING_CHECKPOINTS_API] Global Error: {str(e)}", exc_info=True)
    #         return Response({'error': 'Failed to retrieve upcoming checkpoints.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
   

    @action(detail=False, methods=['get'], url_path='upcoming-checkpoints/(?P<user_id>[^/.]+)')
    def upcoming_checkpoints(self, request, user_id=None):
        try:
            # 1. SETUP: Timezone and expanded Date Range (Yesterday + Today)
            user_tz = get_user_timezone_from_request(request)
            today_user = get_user_today(user_tz)
            yesterday_user = today_user - timedelta(days=1)
            user_now = get_user_now(user_tz)
            user_now_utc = user_now.astimezone(pytz.UTC)

            # Query window for scans: from start of yesterday to end of today
            from patrol_backend.utils.timezone_utils import convert_date_range_to_utc
            start_utc, _ = convert_date_range_to_utc(yesterday_user, yesterday_user, user_tz)
            _, end_utc = convert_date_range_to_utc(today_user, today_user, user_tz)
            end_utc = end_utc + timedelta(days=1)

            # 2. FETCH ASSIGNMENTS
            # Get any assignment that overlaps with yesterday or today
            assignments = list(Assignment.objects.filter(
                start_date__lte=today_user,
                end_date__gte=yesterday_user,
                guard_id=user_id
            ).select_related('shift', 'shift__location'))

            # 3. COLLECT SCANS (Yesterday + Today)
            all_checkins = list(CheckIn.objects.filter(
                guard_id=user_id,
                timestamp__gte=start_utc,
                timestamp__lt=end_utc
            ).order_by('timestamp'))

            used_checkin_ids = set() # To track consumed physical scans
            result = []

            # 4. PROCESS EACH ASSIGNMENT
            for assignment in assignments:
                shift = assignment.shift
                if not shift: continue
                is_overnight = shift.end_time <= shift.start_time
                
                # Check two logical shift cycles: 
                # Instance A (Started Yesterday) and Instance B (Starts Today)
                shift_start_dates = [yesterday_user, today_user]
                
                for shift_start_date in shift_start_dates:
                    # Basic check: is the assignment active on this start date?
                    if not (assignment.start_date <= shift_start_date <= assignment.end_date):
                        continue
                    
                    # --- SMART FILTERING RULES ---
                    # Rule 1: Yesterday's instance must be overnight to still be relevant today
                    if shift_start_date == yesterday_user and not is_overnight:
                        continue
                        
                    # Rule 2: If yesterday's overnight shift ended (e.g., at 06:00 AM) and it's now past that time
                    if shift_start_date == yesterday_user and is_overnight:
                        if user_now.date() == today_user and user_now.time() >= shift.end_time:
                            continue

                    # Rule 3: Hide "tonight's" shift instance if we are currently finishing "last night's" work
                    # This prevents seeing tonight's 21:30 slots while working at 02:00 AM.
                    if shift_start_date == today_user and is_overnight:
                        if user_now.date() == today_user and user_now.time() < shift.end_time:
                            continue

                    # 5. PROCESS CHECKPOINTS
                    for cp in (assignment.checkpoints or []):
                        checkpoint_id_str = cp.get('checkpoint_id')
                        try:
                            checkpoint_id = uuid.UUID(checkpoint_id_str)
                        except (ValueError, TypeError):
                            continue
                        
                        checkpoint_obj = Checkpoint.objects.filter(id=checkpoint_id).first()
                        if not checkpoint_obj: continue

                        checkpoint_time = datetime.strptime(cp['time'], '%H:%M').time()
                        
                        # Calculate the ACTUAL calendar date for this checkpoint instance
                        if is_overnight and checkpoint_time < shift.start_time:
                            # This checkpoint happens after midnight
                            checkpoint_date = shift_start_date + timedelta(days=1)
                        else:
                            # This checkpoint happens on the day the shift started
                            checkpoint_date = shift_start_date

                        # Expected time in UTC for matching and sorting
                        expected_dt_utc = combine_date_time_in_user_tz(checkpoint_date, checkpoint_time, user_tz)

                        # 6. SCAN CONSUMPTION MATCHING (Physical scans only)
                        # This handles multiple scans of the same checkpoint correctly.
                        is_checked_in = False
                        for scan in all_checkins:
                            # Match only physical scans (synced=True) that haven't been consumed
                            if (scan.checkpoint_id == checkpoint_id and
                                getattr(scan, 'synced', True) == True and 
                                scan.id not in used_checkin_ids and 
                                abs(scan.timestamp - expected_dt_utc) <= timedelta(minutes=30)):
                                
                                is_checked_in = True
                                used_checkin_ids.add(scan.id)
                                break

                        # 7. OVERDUE & SYNC LOGIC
                        # Check if 15 minutes have passed since the expected time
                        is_overdue = user_now_utc > (expected_dt_utc + timedelta(minutes=15))

                        if is_checked_in:
                            status, synced = 'completed', True
                        elif is_overdue:
                            status, synced = 'completed', False # status is 'completed' for overdue/missed
                            # Auto-create missed record in DB for reporting
                            CheckIn.objects.get_or_create(
                                guard_id=user_id,
                                shift_id=shift.id,
                                checkpoint_id=checkpoint_id,
                                timestamp=expected_dt_utc,
                                defaults={
                                    'latitude': checkpoint_obj.latitude,
                                    'longitude': checkpoint_obj.longitude,
                                    'synced': False
                                }
                            )
                        else:
                            status, synced = 'pending', None

                        result.append({
                            'assign_id': assignment.id,
                            'checkpoint_id': checkpoint_id,
                            'label': checkpoint_obj.label,
                            'time': expected_dt_utc.astimezone(user_tz).strftime('%H:%M'),
                            'lat': checkpoint_obj.latitude,
                            'lon': checkpoint_obj.longitude,
                            'qr': checkpoint_obj.data,
                            'shift_id': shift.id,
                            'status': status,
                            'synced': synced,
                            'sort_ts': expected_dt_utc.timestamp() # Internal field for chronological sorting
                        })

            # 8. FINAL CHRONOLOGICAL SORTING
            # This ensures that Jan 21 23:30 comes before Jan 22 01:00
            result.sort(key=lambda x: x['sort_ts'])
            
            # Remove internal sort key from final response
            for r in result:
                r.pop('sort_ts', None)
            
            return Response(result)

        except Exception as e:
            logger.error(f"[UPCOMING_CHECKPOINTS_API] Global Error: {str(e)}", exc_info=True)
            return Response({'error': 'Failed to retrieve upcoming checkpoints.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


    @action(detail=False, methods=['get'], url_path='by-guard/(?P<guard_id>[^/.]+)')
    def by_guard(self, request, guard_id=None):
        try:
            assignments = Assignment.objects.filter(guard_id=guard_id, is_deleted=False)
            serializer = self.get_serializer(assignments, many=True)
            return Response(serializer.data)
        except Exception as e:
            logger.error(f"Error fetching assignments by guard: {e}", exc_info=True)
            return Response({'error': 'Failed to retrieve assignments.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


    @action(detail=False, methods=['get'], url_path=r'monthly-location-summary/(?P<location_id>[^/.]+)/(?P<year>\d{4})/(?P<month>\d{1,2})')
    def monthly_location_summary(self, request, location_id=None, year=None, month=None):
        try:
            year = int(year)
            month = int(month)
            num_days = monthrange(year, month)[1]
            
            # Get user timezone - for superadmin viewing location-specific reports, uses location admin's timezone
            user_tz = get_user_timezone_from_request(request, location_id=location_id)
            # Note: This API shows shift assignments (schedules), not attendance records
            # So we always return the full month to show all scheduled shifts, including future dates
            
            # Always generate full month days array (01-01 through 31-01, etc.)
            days = [f"{day:02d}-{month:02d}" for day in range(1, num_days + 1)]

            search = request.query_params.get("search")
            role = request.query_params.get("role")

            # Get all assignments for the location
            assignments = Assignment.objects.select_related('guard', 'shift', 'shift__location').filter(
                shift__location_id=location_id,
                start_date__lte=datetime(year, month, num_days),
                end_date__gte=datetime(year, month, 1),
                is_deleted=False
            )
            if search and str(search).strip():
                s = str(search).strip()
                assignments = assignments.filter(
                    Q(guard__name__icontains=s) | Q(guard__employee_code__icontains=s)
                )
            if role and str(role).strip().lower() not in ('', 'all'):
                assignments = assignments.filter(guard__role__iexact=str(role).strip().lower())

            # Build summary map keyed by guard id (avoids merging duplicate names)
            summary_map = defaultdict(lambda: {
                'name': '',
                'employee_code': '',
                'designation': '',
                'location': '',
                **{day: '-' for day in days}
            })

            for assignment in assignments:
                guard = assignment.guard
                gid = str(guard.id)
                location_name = assignment.shift.location.name
                shift_name = assignment.shift.name

                # Process all days in the month (not limited to today)
                for day in range(1, num_days + 1):
                    current_date = datetime(year, month, day).date()
                    if assignment.start_date <= current_date <= assignment.end_date:
                        key = f"{day:02d}-{month:02d}"
                        summary_map[gid]['name'] = guard.name
                        summary_map[gid]['employee_code'] = getattr(guard, 'employee_code', None) or ''
                        summary_map[gid]['designation'] = (getattr(guard, 'role', None) or '').strip()
                        summary_map[gid]['location'] = location_name
                        summary_map[gid][key] = shift_name

            # Convert to list format
            result = list(summary_map.values())
            return Response({
                'headers': ['Name', 'Emp Code', 'Designation', 'Location'] + days,
                'rows': result
            })

        except Exception as e:
            logger.error(f"Error generating location-based summary: {e}", exc_info=True)
            return Response({'error': 'Failed to generate summary.'}, status=http_status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['get'], url_path=r'monthly-location-summary-excel/(?P<location_id>[^/.]+)/(?P<year>\d{4})/(?P<month>\d{1,2})')
    def monthly_location_summary_excel(self, request, location_id=None, year=None, month=None):
        try:
            year = int(year)
            month = int(month)
            num_days = monthrange(year, month)[1]
            
            # Get user timezone - for superadmin viewing location-specific reports, uses location admin's timezone
            user_tz = get_user_timezone_from_request(request, location_id=location_id)
            # Note: This API shows shift assignments (schedules), not attendance records
            # So we always return the full month to show all scheduled shifts, including future dates
            
            # Always generate full month days array (01-01 through 31-01, etc.)
            days = [f"{day:02d}-{month:02d}" for day in range(1, num_days + 1)]

            search = request.query_params.get("search")
            role = request.query_params.get("role")

            assignments = Assignment.objects.select_related('guard', 'shift', 'shift__location').filter(
                shift__location_id=location_id,
                start_date__lte=datetime(year, month, num_days),
                end_date__gte=datetime(year, month, 1),
                is_deleted=False
            )
            if search and str(search).strip():
                s = str(search).strip()
                assignments = assignments.filter(
                    Q(guard__name__icontains=s) | Q(guard__employee_code__icontains=s)
                )
            if role and str(role).strip().lower() not in ('', 'all'):
                assignments = assignments.filter(guard__role__iexact=str(role).strip().lower())

            summary_map = defaultdict(lambda: {
                'name': '',
                'employee_code': '',
                'designation': '',
                'location': '',
                **{day: '-' for day in days}
            })

            for assignment in assignments:
                guard = assignment.guard
                gid = str(guard.id)
                location_name = assignment.shift.location.name
                shift_name = assignment.shift.name

                # Process all days in the month (not limited to today)
                for day in range(1, num_days + 1):
                    current_date = datetime(year, month, day).date()
                    if assignment.start_date <= current_date <= assignment.end_date:
                        key = f"{day:02d}-{month:02d}"
                        summary_map[gid]['name'] = guard.name
                        summary_map[gid]['employee_code'] = getattr(guard, 'employee_code', None) or ''
                        summary_map[gid]['designation'] = (getattr(guard, 'role', None) or '').strip()
                        summary_map[gid]['location'] = location_name
                        summary_map[gid][key] = shift_name

            # Create Excel workbook
            wb = Workbook()
            ws = wb.active
            ws.title = f"{month:02d}-{year} Summary"

            headers = ['Name', 'Emp Code', 'Designation', 'Location'] + days
            ws.append(headers)

            grouped_rows = defaultdict(list)
            for row in summary_map.values():
                rank = str(row.get('designation') or '').strip().upper() or "UNASSIGNED"
                grouped_rows[rank].append(row)

            for rank in sorted(grouped_rows.keys()):
                for row in sorted(grouped_rows[rank], key=lambda x: str(x.get('name') or '').upper()):
                    ws.append([
                        row['name'],
                        row['employee_code'],
                        row['designation'],
                        row['location'],
                    ] + [row[day] for day in days])

            # Adjust column widths
            for i, column in enumerate(headers, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(12, len(column) + 2)

            # Prepare response
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            filename = f"monthly_summary_{location_id}_{year}_{month}.xlsx"
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            wb.save(response)
            return response

        except Exception as e:
            logger.error(f"Error generating Excel summary: {e}", exc_info=True)
            return Response({'error': 'Failed to generate Excel summary.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class CheckpointViewSet(viewsets.ModelViewSet):
    queryset = Checkpoint.objects.all()
    serializer_class = CheckpointSerializer

    @action(detail=False, methods=['get'], url_path='by-location/(?P<location_id>[^/.]+)')
    def by_location(self, request, location_id=None):
        try:
            checkpoints = Checkpoint.objects.filter(location=location_id, is_deleted=False)
            serializer = self.get_serializer(checkpoints, many=True)
            return Response(serializer.data)
        except Exception as e:
            logger.error(f"Error fetching checkpoints by location: {e}", exc_info=True)
            return Response({'error': 'Failed to retrieve checkpoints.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class SiteSettingViewSet(viewsets.ModelViewSet):
    queryset = SiteSetting.objects.all()
    serializer_class = SiteSettingSerializer

    def get_queryset(self):
        """Simple, fast queryset for retrieving settings."""
        queryset = SiteSetting.objects.filter(is_deleted=False)
        loc_param = self.request.query_params.get('location')
        
        # 1. Superuser logic
        if self.request.user.is_superuser:
            if loc_param:
                if loc_param.lower() == 'null':
                    return queryset.filter(location__isnull=True)
                return queryset.filter(location_id=loc_param)
            
            # For detail actions, don't filter by location so lookup by ID works
            if self.action in ['retrieve', 'update', 'partial_update', 'destroy']:
                return queryset
                
            # Default to Global for Superuser list view
            return queryset.filter(location__isnull=True)

        # 2. Regular users are strictly bound to their location
        if getattr(self.request.user, 'location', None):
            return queryset.filter(location_id=self.request.user.location_id)
        
        # 3. Fallback: Show Global
        return queryset.filter(location__isnull=True)

    def list(self, request, *args, **kwargs):
        """Override list to trigger auto-sync for the current context."""
        user = request.user
        loc_param = request.query_params.get('location')
        
        # Target for sync: 
        # - If superuser and viewing a specific location
        # - If regular user with a location
        target_sync_id = None
        if user.is_superuser:
            if loc_param and loc_param.lower() != 'null':
                target_sync_id = loc_param
        elif getattr(user, 'location', None):
            target_sync_id = user.location_id

        if target_sync_id:
            try:
                self._sync_settings(target_sync_id, user)
            except Exception as e:
                # Log but don't crash the list view
                logger.error(f"[SITE_SETTING_SYNC_ERROR] {e}", exc_info=True)

        return super().list(request, *args, **kwargs)

    def _sync_settings(self, location_id, user):
        """Ensures location has overrides for every global key."""
        # Find global keys that don't have a local counterpart
        # Optimization: Use a subquery to find missing keys
        from django.db.models import Exists, OuterRef
        
        missing_settings = SiteSetting.objects.filter(
            location__isnull=True, 
            is_deleted=False
        ).exclude(
            Exists(
                SiteSetting.objects.filter(
                    key=OuterRef('key'), 
                    location_id=location_id, 
                    is_deleted=False
                )
            )
        )

        to_create = []
        for g_set in missing_settings:
            to_create.append(SiteSetting(
                key=g_set.key,
                value=g_set.value,
                unit=g_set.unit,
                location_id=location_id,
                created_by=user
            ))
        
        if to_create:
            SiteSetting.objects.bulk_create(to_create, ignore_conflicts=True)

    def partial_update(self, request, *args, **kwargs):
        return super().partial_update(request, *args, **kwargs)



    def perform_create(self, serializer):
        user = self.request.user
        # Only Super Admin (no location) can create new keys/settings
        # 1. Permission check: Only Super Admin can create new setting keys
        if not user.is_superuser:
             raise ValidationError("Only Super Admin can create new setting keys.")
        
        # Super admin creations are global (location=NULL)
        serializer.save(created_by=user, location=None)

    def perform_update(self, serializer):
        instance = self.get_object()
        user = self.request.user
        
        # 1. Prevent key name editing
        if 'key' in self.request.data and self.request.data['key'] != instance.key:
            raise ValidationError("Editing the 'key' name is not allowed.")
            
        # 2. Permission check: Location admins can only edit their own settings
        # 2. Permission check: Admin can edit anything, Location admins only their own
        if not user.is_superuser and user.location and instance.location_id != user.location_id:
            raise ValidationError("You do not have permission to edit settings for another location.")
            
        serializer.save(modified_by=user)

    def destroy(self, request, *args, **kwargs):
        user = self.request.user
        instance = self.get_object()
        
        # Only Super Admin can delete settings
        # Only Super Admin can delete settings
        if not user.is_superuser:
            return Response({'error': 'Only Super Admin can delete settings.'}, status=http_status.HTTP_403_FORBIDDEN)
            
        try:
            instance.delete(user=request.user)
            return Response(status=http_status.HTTP_204_NO_CONTENT)
        except Exception as e:
            logger.error(f"Error deleting site setting: {e}", exc_info=True)
            return Response({'error': 'Failed to delete site setting.'}, status=http_status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['post'], url_path='bulk-create')
    def bulk_create(self, request):
        user = self.request.user
        if user.location:
            return Response({'error': 'Only Super Admin can bulk create setting keys.'}, status=http_status.HTTP_403_FORBIDDEN)
            
        try:
            serializer = SiteSettingSerializer(data=request.data, many=True)
            serializer.is_valid(raise_exception=True)
            # Ensure all are global
            serializer.save(created_by=user, location=None)
            return Response(serializer.data, status=http_status.HTTP_201_CREATED)
        except ValidationError as ve:
            return Response({'error': ve.message_dict}, status=http_status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            logger.error(f"Error during bulk create: {e}", exc_info=True)
            return Response({'error': 'Failed to create site settings.'}, status=http_status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['patch'], url_path='bulk-update')
    def bulk_update(self, request):
        user = self.request.user
        data = request.data
        if not isinstance(data, list):
            return Response({'detail': 'Expected a list of objects.'}, status=http_status.HTTP_400_BAD_REQUEST)

        updated_items = []
        try:
            for item in data:
                try:
                    instance = SiteSetting.objects.get(id=item.get('id'))
                except SiteSetting.DoesNotExist:
                    continue
                
                # Permission Check
                if user.location and instance.location_id != user.location_id:
                    continue # Skip settings that don't belong to the user
                
                # Prevent key change
                if 'key' in item and item['key'] != instance.key:
                    return Response({'error': f"Editing the 'key' name for '{instance.key}' is not allowed."}, 
                                   status=http_status.HTTP_400_BAD_REQUEST)

                serializer = SiteSettingSerializer(instance, data=item, partial=True)
                serializer.is_valid(raise_exception=True)
                serializer.save(modified_by=user)
                updated_items.append(serializer.data)

            return Response(updated_items, status=http_status.HTTP_200_OK)
        except Exception as e:
            logger.error(f"Error during bulk update: {e}", exc_info=True)
            return Response({'error': 'Failed to update site settings.'}, status=http_status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['post'], url_path='bulk-delete')
    def bulk_delete(self, request):
        ids = request.data.get('ids', [])
        if not isinstance(ids, list):
           return Response({'detail': 'Expected a list of UUIDs in "ids".'}, status=status.HTTP_400_BAD_REQUEST)


from rest_framework import viewsets, permissions
from rest_framework.response import Response
from rest_framework import status
from .models import CheckpointTemplate, ChecklistItem, ChecklistTemplate
from .serializers import CheckpointTemplateSerializer, ChecklistItemSerializer, ChecklistTemplateSerializer

class CheckpointTemplateViewSet(viewsets.ModelViewSet):

    queryset = CheckpointTemplate.objects.filter(is_deleted=False)
    serializer_class = CheckpointTemplateSerializer
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=False, methods=['get'], url_path='all')
    def get_all_templates(self, request):
        templates = CheckpointTemplate.objects.filter(is_deleted=False)
        result = [self._enrich_template(template) for template in templates]
        return Response(result, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='by-shift-location/(?P<shift_id>[^/.]+)/(?P<location_id>[^/.]+)')
    def by_shift_and_location(self, request, shift_id=None, location_id=None):
        templates = CheckpointTemplate.objects.filter(
            shift_id=shift_id,
            location_id=location_id,
            is_deleted=False
        )
        if not templates.exists():
            return Response({"detail": "No templates found for this shift and location."}, status=status.HTTP_404_NOT_FOUND)

        result = [self._enrich_template(template) for template in templates]
        return Response(result, status=status.HTTP_200_OK)

    @action(detail=False, methods=['get'], url_path='by-location/(?P<location_id>[^/.]+)')
    def by_location(self, request, location_id=None):
        """
        Get all checkpoint templates for a specific location.
        
        URL: /checkpoint-templates/by-location/{location_id}/
        """
        templates = CheckpointTemplate.objects.filter(
            location_id=location_id,
            is_deleted=False
        )
        
        if not templates.exists():
            return Response(
                {"detail": "No templates found for this location."}, 
                status=status.HTTP_404_NOT_FOUND
            )
        
        result = [self._enrich_template(template) for template in templates]
        return Response(result, status=status.HTTP_200_OK)        

    @action(detail=False, methods=['get'], url_path='by-shift/(?P<shift_id>[^/.]+)')
    def by_shift(self, request, shift_id=None):
        templates = CheckpointTemplate.objects.filter(shift_id=shift_id, is_deleted=False)
        if not templates.exists():
            return Response({"detail": "No templates found for this shift."}, status=status.HTTP_404_NOT_FOUND)

        result = []
        for template in templates:
            enriched_checkpoints = []
            for item in template.checkpoints:
                checkpoint_id = item.get("checkpoint_id")
                time = item.get("time")
                checklist_template_id = item.get("checklist_template_id")

                checkpoint = Checkpoint.objects.filter(id=checkpoint_id, is_deleted=False).first()
                checklist_template = None
                if checklist_template_id:
                    checklist_template = ChecklistTemplate.objects.filter(
                        id=checklist_template_id,
                        is_deleted=False
                    ).first()

                if checkpoint:
                    enriched = {
                        "time": time,
                        "checkpoint": {
                            "id": str(checkpoint.id),
                            "label": checkpoint.label,
                            "type": checkpoint.type,
                            "data": checkpoint.data,
                            "latitude": checkpoint.latitude,
                            "longitude": checkpoint.longitude,
                            "location_id": checkpoint.location_id,
                        },
                    }
                    if checklist_template:
                        enriched["checklist_template_id"] = str(checklist_template.id)
                        enriched["checklist_template_name"] = checklist_template.name
                    enriched_checkpoints.append(enriched)

            result.append({
                "template_id": str(template.id),
                "template_name": template.template_name,
                "shift_id": str(template.shift_id),
                "checkpoints": enriched_checkpoints
            })

        return Response(result, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)
    
    def _enrich_template(self, template):
        enriched_checkpoints = []
        for item in template.checkpoints:
            checkpoint_id = item.get("checkpoint_id")
            time = item.get("time")
            checklist_template_id = item.get("checklist_template_id")

            checkpoint = Checkpoint.objects.filter(id=checkpoint_id, is_deleted=False).first()
            checklist_template = None
            if checklist_template_id:
                checklist_template = ChecklistTemplate.objects.filter(
                    id=checklist_template_id,
                    is_deleted=False
                ).first()

            if checkpoint:
                enriched = {
                    "time": time,
                    "checkpoint": {
                        "id": str(checkpoint.id),
                        "label": checkpoint.label,
                        "type": checkpoint.type,
                        "data": checkpoint.data,
                        "latitude": checkpoint.latitude,
                        "longitude": checkpoint.longitude,
                        "location_id": checkpoint.location_id,
                    },
                }
                if checklist_template:
                    enriched["checklist_template_id"] = str(checklist_template.id)
                    enriched["checklist_template_name"] = checklist_template.name
                enriched_checkpoints.append(enriched)

        return {
            "template_id": str(template.id),
            "template_name": template.template_name,
            "shift_id": str(template.shift_id),
            "location_id": str(template.location_id),
            "checkpoints": enriched_checkpoints
        }

    def perform_update(self, serializer):
        serializer.save(modified_by=self.request.user)

    def perform_destroy(self, instance):
        instance.delete(user=self.request.user)


class ChecklistItemViewSet(viewsets.ModelViewSet):
    """
    CRUD for checklist master items.
    """
    serializer_class = ChecklistItemSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        """
        Location-scoped + optional search.
        - Superadmin: can see all or filter by ?location_id=
        - Location admins/others: restricted to their own location only.
        """
        user = self.request.user
        qs = ChecklistItem.objects.filter(is_deleted=False)

        # Restrict by role/location
        user_location_id = getattr(user, 'location_id', None)
        is_super = getattr(user, 'role', None) in ['superadmin', 'super_admin'] or user.is_superuser

        param_location_id = self.request.query_params.get('location_id')

        if is_super:
            if param_location_id and param_location_id != 'All':
                qs = qs.filter(location_id=param_location_id)
        else:
            # Non-super users only see their own location's items
            if user_location_id:
                qs = qs.filter(location_id=user_location_id)
            else:
                qs = qs.none()

        # Simple search by label
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(label__icontains=search)

        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    def perform_update(self, serializer):
        serializer.save(modified_by=self.request.user)

    def perform_destroy(self, instance):
        instance.delete(user=self.request.user)


class ChecklistTemplateViewSet(viewsets.ModelViewSet):
    """
    CRUD for checklist templates (groups of items).
    """
    serializer_class = ChecklistTemplateSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        """
        Location-scoped + optional search.
        - Superadmin: can see all or filter by ?location_id=
        - Location admins/others: restricted to their own location only.
        """
        user = self.request.user
        qs = ChecklistTemplate.objects.filter(is_deleted=False)

        user_location_id = getattr(user, 'location_id', None)
        is_super = getattr(user, 'role', None) in ['superadmin', 'super_admin'] or user.is_superuser

        param_location_id = self.request.query_params.get('location_id')

        if is_super:
            if param_location_id and param_location_id != 'All':
                qs = qs.filter(location_id=param_location_id)
        else:
            if user_location_id:
                qs = qs.filter(location_id=user_location_id)
            else:
                qs = qs.none()

        # Simple search by template name
        search = self.request.query_params.get('search')
        if search:
            qs = qs.filter(name__icontains=search)

        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    def perform_update(self, serializer):
        serializer.save(modified_by=self.request.user)

    def perform_destroy(self, instance):
        instance.delete(user=self.request.user)

# from rest_framework.views import APIView
# from rest_framework.response import Response
# from rest_framework import status, permissions
# from .models import CheckpointTemplate, Checkpoint
# from .serializers import CheckpointTemplateSerializer
# from django.shortcuts import get_object_or_404

# class CheckpointTemplateByShiftView(APIView):
#     permission_classes = [permissions.IsAuthenticated]

#     def get(self, request, shift_id):
#         templates = CheckpointTemplate.objects.filter(shift_id=shift_id, is_deleted=False)
#         if not templates.exists():
#             return Response({"detail": "No templates found for this shift."}, status=status.HTTP_404_NOT_FOUND)

#         result = []
#         for template in templates:
#             enriched_checkpoints = []
#             for item in template.checkpoints:
#                 checkpoint_id = item.get("checkpoint_id")
#                 time = item.get("time")
#                 checkpoint = Checkpoint.objects.filter(id=checkpoint_id, is_deleted=False).first()
#                 if checkpoint:
#                     enriched_checkpoints.append({
#                         "time": time,
#                         "checkpoint": {
#                             "id": str(checkpoint.id),
#                             "label": checkpoint.label,
#                             "type": checkpoint.type,
#                             "data": checkpoint.data,
#                             "latitude": checkpoint.latitude,
#                             "longitude": checkpoint.longitude,
#                             "location_id": checkpoint.location_id,
#                         }
#                     })
#             result.append({
#                 "template_id": str(template.id),
#                 "template_name": template.template_name,
#                 "shift_id": str(template.shift_id),
#                 "checkpoints": enriched_checkpoints
#             })

#         return Response(result, status=status.HTTP_200_OK)