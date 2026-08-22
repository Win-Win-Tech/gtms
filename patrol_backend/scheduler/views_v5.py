"""v5 assignment views — copies of live AssignmentViewSet with posted site. Live URLs unchanged."""

from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime

from django.http import HttpResponse
from openpyxl import Workbook
from rest_framework import permissions, status
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response

from authapp.site_access import (
    allowed_site_ids,
    assert_caller_can_access_site,
    get_site_or_error,
    is_org_admin,
)
from scheduler.daily_site import clear_daily_sites
from scheduler.models import Assignment, AssignmentDailySite
from scheduler.serializers_v5 import AssignmentSerializerV5
from scheduler.views import AssignmentViewSet, _get_monthly_location_summary_data_v2
from patrol_backend.utils.timezone_utils import get_user_today, get_user_timezone_from_request


def _v5_error(exc):
    if isinstance(exc, ValidationError):
        detail = exc.detail
        if isinstance(detail, dict):
            return Response(detail, status=status.HTTP_400_BAD_REQUEST)
        return Response({"error": detail}, status=status.HTTP_400_BAD_REQUEST)
    if isinstance(exc, PermissionDenied):
        return Response({"error": str(exc)}, status=status.HTTP_403_FORBIDDEN)
    raise exc


def _monthly_summary_site_scope_kwargs(request, location_id, year, month):
    raw = request.query_params.get("site_id")
    site_id = None if not raw or str(raw).lower() in ("all", "null", "undefined") else raw
    caller = request.user
    year_i, month_i = int(year), int(month)
    m_start = date(year_i, month_i, 1)
    m_end = date(year_i, month_i, monthrange(year_i, month_i)[1])

    if site_id:
        site = get_site_or_error(site_id)
        assert_caller_can_access_site(caller, site)
        if str(site.location_id) != str(location_id):
            raise ValidationError({"site_id": "Site does not belong to that organisation."})
        guard_ids = list(
            AssignmentDailySite.objects.filter(
                site_id=site_id,
                date__gte=m_start,
                date__lte=m_end,
            ).values_list("guard_id", flat=True).distinct()
        )
        return {"site_id": str(site_id), "site_guard_ids": guard_ids}

    if caller.is_superuser or is_org_admin(caller) or getattr(caller, "all_org_sites", False):
        return {}

    allowed = allowed_site_ids(caller)
    guard_ids = list(
        AssignmentDailySite.objects.filter(
            site_id__in=allowed,
            date__gte=m_start,
            date__lte=m_end,
        ).values_list("guard_id", flat=True).distinct()
    )
    return {"site_guard_ids": guard_ids}


class AssignmentViewSetV5(AssignmentViewSet):
    serializer_class = AssignmentSerializerV5
    permission_classes = [permissions.IsAuthenticated]

    def _posted_on_date(self):
        raw = self.request.query_params.get("date")
        if raw:
            try:
                return datetime.strptime(raw, "%Y-%m-%d").date()
            except ValueError:
                raise ValidationError({"date": "Use YYYY-MM-DD."})
        tz = get_user_timezone_from_request(self.request)
        return get_user_today(tz)

    def _filter_by_posted_site(self, qs):
        site_id = self.request.query_params.get("site_id")
        if not site_id:
            return qs
        site = get_site_or_error(site_id)
        assert_caller_can_access_site(self.request.user, site)
        on_date = self._posted_on_date()
        guard_ids = AssignmentDailySite.objects.filter(
            site_id=site.id,
            date=on_date,
        ).values_list("guard_id", flat=True)
        return qs.filter(guard_id__in=guard_ids)

    def get_queryset(self):
        return self._filter_by_posted_site(
            super().get_queryset().prefetch_related("daily_sites__site")
        )

    def list(self, request, *args, **kwargs):
        try:
            return super().list(request, *args, **kwargs)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)
        except PermissionDenied as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)

    def create(self, request, *args, **kwargs):
        try:
            return super().create(request, *args, **kwargs)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)

    def update(self, request, *args, **kwargs):
        try:
            return super().update(request, *args, **kwargs)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        clear_daily_sites(instance)
        return super().destroy(request, *args, **kwargs)

    @action(detail=False, methods=["get"], url_path="by-guard/(?P<guard_id>[^/.]+)")
    def by_guard(self, request, guard_id=None):
        try:
            try:
                limit = min(max(int(request.query_params.get("limit") or 25), 1), 250)
            except (TypeError, ValueError):
                limit = 25
            try:
                offset = max(int(request.query_params.get("offset") or 0), 0)
            except (TypeError, ValueError):
                offset = 0

            qs = Assignment.objects.filter(guard_id=guard_id, is_deleted=False).prefetch_related(
                "daily_sites__site"
            )
            qs = self._filter_by_posted_site(qs).order_by("-start_date", "-end_date", "-created_on")
            total = qs.count()
            page = qs[offset : offset + limit]
            serializer = self.get_serializer(page, many=True)
            return Response({
                "count": total,
                "limit": limit,
                "offset": offset,
                "results": serializer.data,
            })
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)
        except PermissionDenied as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)

    @action(detail=False, methods=["post"], url_path="bulk-create")
    def bulk_create(self, request):
        assignments_data = request.data
        if not isinstance(assignments_data, list):
            return Response({"error": "Expected a list of assignments"}, status=status.HTTP_400_BAD_REQUEST)

        created = []
        errors = []
        for idx, data in enumerate(assignments_data):
            serializer = self.get_serializer(data=data)
            if serializer.is_valid():
                try:
                    serializer.save()
                    created.append(serializer.data)
                except ValidationError as e:
                    errors.append({"index": idx, "errors": e.detail, "data": data})
                except PermissionDenied as e:
                    errors.append({"index": idx, "errors": {"site_id": [str(e)]}, "data": data})
            else:
                errors.append({"index": idx, "errors": serializer.errors, "data": data})

        if errors:
            return Response({"created": created, "errors": errors}, status=status.HTTP_207_MULTI_STATUS)
        return Response({"created": created}, status=status.HTTP_201_CREATED)

    @action(
        detail=False,
        methods=["get"],
        url_path=r"v5/monthly-location-summary/(?P<location_id>[^/.]+)/(?P<year>\d{4})/(?P<month>\d{1,2})",
    )
    def monthly_location_summary_v5(self, request, location_id=None, year=None, month=None):
        try:
            scope = _monthly_summary_site_scope_kwargs(request, location_id, year, month)
            user_id = (
                request.query_params.get("user_id")
                or request.query_params.get("id")
                or request.query_params.get("guard")
            )
            res = _get_monthly_location_summary_data_v2(
                location_id=location_id,
                year=year,
                month=month,
                search=request.query_params.get("search"),
                role=request.query_params.get("role"),
                request=request,
                user_id=user_id,
                site_id=scope.get("site_id"),
                site_guard_ids=scope.get("site_guard_ids"),
            )
            return Response({"headers": res["headers"], "rows": res["rows"]})
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)
        except Exception as e:
            import logging

            logging.getLogger(__name__).error("monthly_location_summary_v5: %s", e, exc_info=True)
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(
        detail=False,
        methods=["get"],
        url_path=r"v5/monthly-location-summary-excel/(?P<location_id>[^/.]+)/(?P<year>\d{4})/(?P<month>\d{1,2})",
    )
    def monthly_location_summary_excel_v5(self, request, location_id=None, year=None, month=None):
        try:
            scope = _monthly_summary_site_scope_kwargs(request, location_id, year, month)
            user_id = (
                request.query_params.get("user_id")
                or request.query_params.get("id")
                or request.query_params.get("guard")
            )
            res = _get_monthly_location_summary_data_v2(
                location_id=location_id,
                year=year,
                month=month,
                search=request.query_params.get("search"),
                role=request.query_params.get("role"),
                request=request,
                user_id=user_id,
                site_id=scope.get("site_id"),
                site_guard_ids=scope.get("site_guard_ids"),
            )

            wb = Workbook()
            ws = wb.active
            ws.title = f"{int(month):02d}-{year} Summary v5"
            ws.append(res["headers"])

            grouped_rows = defaultdict(list)
            for row in res["rows"]:
                rank = str(row.get("designation") or "").strip().upper() or "UNASSIGNED"
                grouped_rows[rank].append(row)

            for rank in sorted(grouped_rows.keys()):
                for row in sorted(grouped_rows[rank], key=lambda x: str(x.get("name") or "").upper()):
                    ws.append(
                        [
                            row["name"],
                            row["employee_code"],
                            row["designation"],
                            row["location"],
                        ]
                        + [row[day] for day in res["days"]]
                    )

            from openpyxl.utils import get_column_letter

            for i, column in enumerate(res["headers"], 1):
                ws.column_dimensions[get_column_letter(i)].width = max(12, len(column) + 2)

            response = HttpResponse(
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            filename = f"monthly_summary_v5_{location_id}_{year}_{month}.xlsx"
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            wb.save(response)
            return response
        except (ValidationError, PermissionDenied) as exc:
            return _v5_error(exc)
        except Exception as e:
            import logging

            logging.getLogger(__name__).error("monthly_location_summary_excel_v5: %s", e, exc_info=True)
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
