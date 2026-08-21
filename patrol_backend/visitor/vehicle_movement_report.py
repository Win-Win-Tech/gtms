"""Vehicle movement summary: check-in (in) / check-out (out) counts by vehicle type."""
from datetime import datetime, time as dt_time, timedelta
from io import BytesIO

import pytz
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone as dj_timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from patrol_backend.utils.timezone_utils import (
    get_user_timezone_from_request,
    get_user_today,
    to_user_timezone,
)

from .models import VisitorEntry

# Display order for report rows (uses existing vehicle_type choices only)
VEHICLE_REPORT_ORDER = [
    "car",
    "truck",
    "van",
    "motorcycle",
    "bus",
    "other",
]


def _vehicle_label(code):
    for value, label in VisitorEntry.VEHICLE_TYPE_CHOICES:
        if value == code:
            return label
    return (code or "").replace("_", " ").title() or "—"


def _parse_report_date(value):
    if not value:
        return None
    text = str(value).strip()[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None


def _parse_hhmm(value, default):
    if not value:
        return default
    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            parsed = datetime.strptime(text, fmt).time()
            return dt_time(parsed.hour, parsed.minute)
        except ValueError:
            continue
    return default


def _fmt_time_display(t):
    if not t:
        return ""
    return t.strftime("%H:%M")


def _local_window_to_utc(report_date, start_t, end_t, user_tz):
    return _local_range_to_utc(report_date, start_t, report_date, end_t, user_tz)


def _local_range_to_utc(start_date, start_t, end_date, end_t, user_tz):
    start_local = user_tz.localize(datetime.combine(start_date, start_t))
    end_local = user_tz.localize(datetime.combine(end_date, end_t))
    if end_t.hour == 23 and end_t.minute == 59:
        end_local = end_local.replace(second=59, microsecond=999999)
    if end_local <= start_local:
        raise ValueError("End must be after start")
    return start_local.astimezone(pytz.UTC), end_local.astimezone(pytz.UTC)


def resolve_vehicle_movement_window(
    request,
    location_id,
    date_filter,
    start_date_str=None,
    end_date_str=None,
    start_time_str=None,
    end_time_str=None,
):
    """Resolve UTC window and display labels from date_filter (matches visitor list filters)."""
    user_tz = get_user_timezone_from_request(request, location_id=location_id)
    user_today = get_user_today(user_tz)
    date_filter = (date_filter or "today").lower()
    full_day_start = dt_time(0, 0)
    full_day_end = dt_time(23, 59)

    if date_filter == "all":
        return {
            "start_utc": None,
            "end_utc": None,
            "date_filter": date_filter,
            "date_range_display": "All dates",
            "time_from": "",
            "time_to": "",
            "show_time": False,
        }

    if date_filter == "custom":
        if not start_date_str or not end_date_str:
            raise ValueError("start_date and end_date are required for custom date filter")
        start_d = _parse_report_date(start_date_str)
        end_d = _parse_report_date(end_date_str)
        if not start_d or not end_d:
            raise ValueError("Invalid date format. Use YYYY-MM-DD.")
        start_t = _parse_hhmm(start_time_str, dt_time(0, 0))
        end_t = _parse_hhmm(end_time_str, dt_time(23, 59))
        show_time = True
    elif date_filter == "today":
        start_d = end_d = user_today
        start_t = full_day_start
        end_t = full_day_end
        show_time = False
    elif date_filter in ("upcoming", "future"):
        start_d = user_today + timedelta(days=1)
        end_d = start_d + timedelta(days=365)
        start_t = full_day_start
        end_t = full_day_end
        show_time = False
    elif date_filter in ("week", "this_week"):
        start_d = user_today - timedelta(days=user_today.weekday())
        end_d = start_d + timedelta(days=6)
        start_t = full_day_start
        end_t = full_day_end
        show_time = False
    elif date_filter in ("month", "this_month"):
        start_d = user_today.replace(day=1)
        if user_today.month == 12:
            end_d = user_today.replace(year=user_today.year + 1, month=1, day=1) - timedelta(
                days=1
            )
        else:
            end_d = user_today.replace(month=user_today.month + 1, day=1) - timedelta(days=1)
        start_t = full_day_start
        end_t = full_day_end
        show_time = False
    else:
        start_d = end_d = user_today
        start_t = full_day_start
        end_t = full_day_end
        show_time = False

    start_utc, end_utc = _local_range_to_utc(start_d, start_t, end_d, end_t, user_tz)
    if start_d == end_d:
        date_range_display = start_d.strftime("%d/%m/%Y")
    else:
        date_range_display = (
            f"{start_d.strftime('%d/%m/%Y')} – {end_d.strftime('%d/%m/%Y')}"
        )

    time_from = _fmt_time_display(start_t) if show_time else ""
    time_to = _fmt_time_display(end_t) if show_time else ""

    return {
        "start_utc": start_utc,
        "end_utc": end_utc,
        "date_filter": date_filter,
        "date_range_display": date_range_display,
        "time_from": time_from,
        "time_to": time_to,
        "show_time": show_time,
        "start_date": start_d.isoformat(),
        "end_date": end_d.isoformat(),
    }


def _vehicle_base_qs(location_id, site_id=None):
    qs = (
        VisitorEntry.objects.filter(is_deleted=False, location_id=location_id)
        .exclude(status__in=[VisitorEntry.STATUS_CANCELLED, VisitorEntry.STATUS_REVERTED])
        .exclude(Q(vehicle_type__isnull=True) | Q(vehicle_type__exact=""))
    )
    if site_id:
        qs = qs.filter(site_id=site_id)
    return qs


def build_vehicle_movement_report(
    request,
    location_id,
    date_filter="today",
    start_date_str=None,
    end_date_str=None,
    start_time_str=None,
    end_time_str=None,
    site_id=None,
):
    """
    In = check_in_time in window; Out = check_out_time in window.
    Grouped by vehicle_type for one location + date/time range.
    Optional site_id scopes to that site only.
    """
    from scheduler.models import Location

    window = resolve_vehicle_movement_window(
        request,
        location_id,
        date_filter,
        start_date_str,
        end_date_str,
        start_time_str,
        end_time_str,
    )
    start_utc = window["start_utc"]
    end_utc = window["end_utc"]

    loc = Location.objects.filter(id=location_id, is_deleted=False).first()
    location_name = loc.name if loc else "—"

    base = _vehicle_base_qs(location_id, site_id=site_id)

    in_filters = Q()
    out_filters = Q()
    if start_utc and end_utc:
        in_filters = Q(check_in_time__gte=start_utc, check_in_time__lte=end_utc)
        out_filters = Q(check_out_time__gte=start_utc, check_out_time__lte=end_utc)
    else:
        in_filters = Q(check_in_time__isnull=False)
        out_filters = Q(check_out_time__isnull=False)

    in_map = {
        row["vehicle_type"]: row["count"]
        for row in base.filter(in_filters)
        .values("vehicle_type")
        .annotate(count=Count("id"))
    }

    out_map = {
        row["vehicle_type"]: row["count"]
        for row in base.filter(out_filters)
        .values("vehicle_type")
        .annotate(count=Count("id"))
    }

    all_types = set(in_map) | set(out_map)

    rows = []
    total_in = 0
    total_out = 0
    for code in VEHICLE_REPORT_ORDER:
        count_in = in_map.get(code, 0)
        count_out = out_map.get(code, 0)
        label = _vehicle_label(code)
        rows.append(
            {
                "vehicle_type": code,
                "label": label,
                "in": count_in,
                "out": count_out,
            }
        )
        total_in += count_in
        total_out += count_out
    # Any extra types not in the default order
    for code in sorted(all_types):
        if code in VEHICLE_REPORT_ORDER:
            continue
        count_in = in_map.get(code, 0)
        count_out = out_map.get(code, 0)
        rows.append(
            {
                "vehicle_type": code,
                "label": _vehicle_label(code),
                "in": count_in,
                "out": count_out,
            }
        )
        total_in += count_in
        total_out += count_out

    return {
        "location_id": str(location_id),
        "location_name": location_name,
        "date_filter": window["date_filter"],
        "report_date_display": window["date_range_display"],
        "date_range_display": window["date_range_display"],
        "time_from": window["time_from"],
        "time_to": window["time_to"],
        "show_time": window["show_time"],
        "start_date": window.get("start_date"),
        "end_date": window.get("end_date"),
        "rows": rows,
        "total_in": total_in,
        "total_out": total_out,
    }


MOVEMENT_HEADERS = ["Vehicle Type", "In", "Out"]
PDF_FOOTER_TEXT = "Generated by CloudGen Technologies"


def _fmt_report_date(d):
    if not d:
        return ""
    if isinstance(d, str):
        parsed = _parse_report_date(d)
        if not parsed:
            return d
        d = parsed
    return d.strftime("%b %d %Y")


def generate_vehicle_movement_excel(report_data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicle Movement"

    ws.append(MOVEMENT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in report_data.get("rows") or []:
        ws.append([row.get("label", ""), row.get("in", 0), row.get("out", 0)])

    ws.append(["TOTAL", report_data.get("total_in", 0), report_data.get("total_out", 0)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_slug = (report_data.get("start_date") or "report").replace("-", "")
    filename = f"vehicle_movement_{date_slug}.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _draw_vehicle_movement_pdf_header_and_footer(canvas, doc):
    """Dashboard/visitor-style repeating header, header line, and footer on every page."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, Table, TableStyle

    canvas.saveState()

    canvas.setFont("Helvetica-Oblique", 8)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawCentredString(
        doc.leftMargin + doc.width / 2.0,
        0.15 * inch,
        PDF_FOOTER_TEXT,
    )
    canvas.drawRightString(
        doc.leftMargin + doc.width,
        0.15 * inch,
        f"Page {canvas.getPageNumber()}",
    )

    org_name = getattr(doc, "org_name", "—")
    dept_label = getattr(doc, "dept_label", "—")
    range_start = getattr(doc, "range_start", None)
    range_end = getattr(doc, "range_end", None)
    printed_at = getattr(doc, "printed_at", None)
    time_from = getattr(doc, "time_from", "")
    time_to = getattr(doc, "time_to", "")
    show_time = getattr(doc, "show_time", False)

    styles = getSampleStyleSheet()

    def _para_style(name, **kwargs):
        base = dict(
            leftIndent=0,
            rightIndent=0,
            firstLineIndent=0,
            spaceBefore=0,
            spaceAfter=0,
        )
        base.update(kwargs)
        return ParagraphStyle(name, **base)

    title_style = _para_style(
        "VmHeaderTitle",
        parent=styles["Normal"],
        fontSize=11,
        fontName="Helvetica-Bold",
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    date_style = _para_style(
        "VmHeaderDate",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    meta_style = _para_style(
        "VmHeaderMeta",
        parent=styles["Normal"],
        fontSize=9,
        fontName="Helvetica-Bold",
        alignment=TA_LEFT,
        leading=11,
    )
    meta_right_style = _para_style(
        "VmHeaderMetaRight",
        parent=meta_style,
        alignment=TA_RIGHT,
    )

    title_p = Paragraph("Vehicle Movement Report", title_style)
    if range_start and range_end:
        date_line = f"{_fmt_report_date(range_start)} To {_fmt_report_date(range_end)}"
        if show_time and time_from and time_to:
            date_line = f"{date_line}  ({time_from} – {time_to})"
    else:
        date_line = getattr(doc, "date_range_display", "") or "—"
    date_p = Paragraph(date_line, date_style)
    printed_text = ""
    if printed_at:
        printed_text = f"Printed On : {printed_at.strftime('%b %d %Y %H:%M')}"

    header_meta = Table(
        [
            [
                Paragraph(f"Company &nbsp; {org_name}", meta_style),
                Paragraph(printed_text, meta_right_style),
            ],
            [Paragraph(f"Site &nbsp; {dept_label}", meta_style), ""],
        ],
        colWidths=[doc.width * 0.55, doc.width * 0.45],
    )
    header_meta.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    header_meta.hAlign = "LEFT"

    top_y = doc.height + doc.bottomMargin + doc.topMargin
    title_p.wrap(doc.width, doc.topMargin)
    title_p.drawOn(canvas, doc.leftMargin, top_y - 18)
    date_p.wrap(doc.width, doc.topMargin)
    date_p.drawOn(canvas, doc.leftMargin, top_y - 32)
    header_meta.wrap(doc.width, doc.topMargin)
    header_meta.drawOn(canvas, doc.leftMargin, top_y - 75)

    canvas.setStrokeColor(colors.HexColor("#cbd5e1"))
    canvas.setLineWidth(0.75)
    canvas.line(doc.leftMargin, top_y - 80, doc.leftMargin + doc.width, top_y - 80)

    canvas.restoreState()


def generate_vehicle_movement_pdf(report_data, request=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    org_name = report_data.get("location_name") or "—"
    printed_at = None
    if request is not None:
        user_tz = get_user_timezone_from_request(
            request, location_id=report_data.get("location_id")
        )
        printed_at = to_user_timezone(dj_timezone.now(), user_tz)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=1.2 * inch,
        bottomMargin=0.4 * inch,
    )
    doc.org_name = org_name
    doc.dept_label = org_name
    doc.range_start = report_data.get("start_date")
    doc.range_end = report_data.get("end_date")
    doc.printed_at = printed_at
    doc.time_from = report_data.get("time_from") or ""
    doc.time_to = report_data.get("time_to") or ""
    doc.show_time = bool(report_data.get("show_time"))
    doc.date_range_display = report_data.get("date_range_display") or ""

    styles = getSampleStyleSheet()

    def _para_style(name, **kwargs):
        base = dict(
            leftIndent=0,
            rightIndent=0,
            firstLineIndent=0,
            spaceBefore=0,
            spaceAfter=0,
        )
        base.update(kwargs)
        return ParagraphStyle(name, **base)

    cell_style = _para_style(
        "VmCell",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_LEFT,
    )
    header_style = _para_style(
        "VmHeader",
        parent=styles["Normal"],
        fontSize=9,
        fontName="Helvetica-Bold",
        alignment=TA_CENTER,
    )

    data = [[Paragraph(h, header_style) for h in MOVEMENT_HEADERS]]
    for row in report_data.get("rows") or []:
        data.append(
            [
                Paragraph(str(row.get("label", "")), cell_style),
                Paragraph(str(row.get("in", 0)), cell_style),
                Paragraph(str(row.get("out", 0)), cell_style),
            ]
        )
    data.append(
        [
            Paragraph("<b>TOTAL</b>", cell_style),
            Paragraph(f"<b>{report_data.get('total_in', 0)}</b>", cell_style),
            Paragraph(f"<b>{report_data.get('total_out', 0)}</b>", cell_style),
        ]
    )

    col_widths = [doc.width * 0.55, doc.width * 0.225, doc.width * 0.225]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )

    doc.build(
        [table],
        onFirstPage=_draw_vehicle_movement_pdf_header_and_footer,
        onLaterPages=_draw_vehicle_movement_pdf_header_and_footer,
    )
    buffer.seek(0)
    date_slug = (report_data.get("start_date") or "report").replace("-", "")
    filename = f"vehicle_movement_{date_slug}.pdf"
    return buffer.getvalue(), filename
