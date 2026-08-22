"""Quick Pay Report Excel export — hourly and monthly sections."""
from __future__ import annotations

import os
from io import BytesIO

from django.conf import settings
from django.utils import timezone


HOURLY_COLUMNS = [
    ("Employee Code", "employee_code"),
    ("Date", "date_label"),
    ("Name", "name"),
    ("Rate", "rate"),
    ("Worked Duration (hours)", "worked_hours"),
    ("Gross", "gross_earnings"),
    ("PF", "pf"),
    ("ESI", "esi"),
    ("Advance Recovery", "advance_recovery"),
    ("Net Paid", "net_paid"),
]

MONTHLY_COLUMNS = [
    ("Employee Code", "employee_code"),
    ("Date", "date_label"),
    ("Name", "name"),
    ("Gross", "gross_salary"),
    ("Paid Days", "paid_days"),
    ("PF", "pf"),
    ("ESI", "esi"),
    ("Advance Recovery", "advance_recovery"),
    ("Net Paid", "net_paid"),
]


def _write_section(ws, start_row, section_title, columns, rows, styles):
    title_font, header_font, cell_font, border, center, left = styles
    num_cols = len(columns)

    ws.merge_cells(
        start_row=start_row, start_column=1, end_row=start_row, end_column=num_cols
    )
    ws.cell(row=start_row, column=1, value=section_title)
    ws.cell(row=start_row, column=1).font = title_font
    ws.cell(row=start_row, column=1).alignment = center

    header_row = start_row + 1
    for col_idx, (label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = header_font
        cell.border = border
        cell.alignment = center

    data_start = header_row + 1
    numeric_keys = {
        "rate", "worked_hours", "salary", "gross_salary", "gross_earnings",
        "paid_days", "pf", "esi", "pt", "advance_recovery", "net_paid",
    }
    left_keys = {"name", "employee_code", "date_label"}

    if not rows:
        ws.merge_cells(
            start_row=data_start,
            start_column=1,
            end_row=data_start,
            end_column=num_cols,
        )
        cell = ws.cell(row=data_start, column=1, value="No employees in this section.")
        cell.font = cell_font
        cell.border = border
        cell.alignment = left
        return data_start + 1

    for row_idx, row in enumerate(rows, start=data_start):
        for col_idx, (_, key) in enumerate(columns, start=1):
            val = row.get(key, "")
            if key in numeric_keys:
                val = float(val) if val is not None else 0
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.border = border
            if key in numeric_keys:
                cell.alignment = center
                cell.number_format = "0.00" if key != "paid_days" else "0.##"
            else:
                cell.alignment = left if key in left_keys else center

    return data_start + len(rows)


def build_hourly_wage_summary_excel_bytes(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Quick Pay Report"

    org_name = context.get("org_name") or ""
    site_name = context.get("site_name") or ""
    period_label = context.get("period_label") or ""
    hourly_rows = context.get("hourly_rows") or []
    monthly_rows = context.get("monthly_rows") or []

    title_font = Font(name="Calibri", bold=True, size=12)
    section_font = Font(name="Calibri", bold=True, size=11)
    header_font = Font(name="Calibri", bold=True, size=10)
    cell_font = Font(name="Calibri", bold=True, size=10)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    styles = (section_font, header_font, cell_font, border, center, left)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    title_parts = ["Quick Pay Report", org_name, site_name, period_label]
    ws.cell(row=1, column=1, value=" — ".join(part for part in title_parts if part))
    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=1, column=1).alignment = center

    next_row = 3
    next_row = _write_section(
        ws, next_row, "Hourly Employees", HOURLY_COLUMNS, hourly_rows, styles
    )
    next_row += 1
    _write_section(
        ws, next_row, "Monthly Employees", MONTHLY_COLUMNS, monthly_rows, styles
    )

    widths = [14, 24, 28, 12, 18, 12, 10, 10, 16, 12]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_hourly_wage_summary_excel_internal(context):
    xlsx_bytes = build_hourly_wage_summary_excel_bytes(context)
    filename = f"quick_pay_report_{timezone.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    file_path = os.path.join(settings.MEDIA_ROOT, filename)
    os.makedirs(settings.MEDIA_ROOT, exist_ok=True)
    with open(file_path, "wb") as fh:
        fh.write(xlsx_bytes)
    row_count = len(context.get("hourly_rows") or []) + len(context.get("monthly_rows") or [])
    return {
        "file_path": file_path,
        "filename": filename,
        "row_count": row_count,
        "xlsx_bytes": xlsx_bytes,
    }
