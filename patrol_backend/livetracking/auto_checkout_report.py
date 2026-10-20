"""
Auto-checkout reporting service: query filtering, dashboard serialization,
Excel (.xlsx) generation with openpyxl, and PDF (.pdf) generation with ReportLab.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, time
from typing import Optional

from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dashboard.models import CheckInLog
from patrol_backend.utils.timezone_utils import (
    from_user_timezone,
    get_user_timezone,
    to_user_timezone,
)
from .ws_groups import _is_superadmin

logger = logging.getLogger(__name__)


def filter_auto_checkout_queryset(request, user):
    """
    Filter CheckInLog rows where type='checkout' and is_auto_checkout=True.
    Applies organization scoping and user request query parameters.
    """
    qs = (
        CheckInLog.objects.filter(type="checkout", is_auto_checkout=True)
        .select_related("guard", "site", "org_location", "shift", "assignment")
        .order_by("-timestamp")
    )

    is_super = _is_superadmin(user)
    user_location_id = getattr(user, "location_id", None)

    # Org scoping for non-superadmin
    if not is_super:
        if user_location_id:
            qs = qs.filter(org_location_id=user_location_id)
        else:
            return qs.none()

    # Query param filters
    location_id = request.query_params.get("location_id")
    if is_super and location_id and location_id not in ("All", "all", "null"):
        qs = qs.filter(org_location_id=location_id)

    site_id = request.query_params.get("site_id")
    if site_id and site_id not in ("All", "all", "null"):
        qs = qs.filter(site_id=site_id)

    guard_id = (
        request.query_params.get("guard_id")
        or request.query_params.get("user_id")
        or request.query_params.get("user")
    )
    if guard_id and guard_id not in ("All", "all", "null"):
        qs = qs.filter(guard_id=guard_id)

    role = request.query_params.get("role")
    if role and role not in ("All", "all", "null"):
        qs = qs.filter(guard__role__iexact=role)

    search = request.query_params.get("search", "").strip()
    if search:
        qs = qs.filter(
            Q(guard__name__icontains=search)
            | Q(guard__email__icontains=search)
            | Q(guard__phone__icontains=search)
            | Q(guard__employee_code__icontains=search)
            | Q(site__name__icontains=search)
            | Q(auto_checkout_reason__icontains=search)
        )

    # Date range filtering
    start_date_str = request.query_params.get("start_date")
    end_date_str = request.query_params.get("end_date")
    user_tz = get_user_timezone(user)

    if start_date_str:
        try:
            start_d = datetime.strptime(start_date_str.strip(), "%Y-%m-%d").date()
            start_dt = from_user_timezone(datetime.combine(start_d, time.min), user_tz)
            qs = qs.filter(timestamp__gte=start_dt)
        except (ValueError, TypeError):
            pass

    if end_date_str:
        try:
            end_d = datetime.strptime(end_date_str.strip(), "%Y-%m-%d").date()
            end_dt = from_user_timezone(datetime.combine(end_d, time.max), user_tz)
            qs = qs.filter(timestamp__lte=end_dt)
        except (ValueError, TypeError):
            pass

    return qs


def serialize_auto_checkout_log(log: CheckInLog, requesting_user) -> dict:
    """Format CheckInLog into JSON dictionary for web dashboard."""
    guard = log.guard
    site = log.site
    org_loc = log.org_location
    shift = log.shift

    local_dt = to_user_timezone(log.timestamp, requesting_user)
    timestamp_str = local_dt.strftime("%Y-%m-%d %H:%M:%S") if local_dt else ""

    return {
        "id": log.id,
        "timestamp": timestamp_str,
        "timestamp_iso": local_dt.isoformat() if local_dt else "",
        "guard_id": str(guard.id) if guard else None,
        "guard_name": getattr(guard, "name", None) or getattr(guard, "email", "Unknown"),
        "employee_code": getattr(guard, "employee_code", "") or "",
        "guard_email": getattr(guard, "email", ""),
        "guard_phone": getattr(guard, "phone", "") or "",
        "role": getattr(guard, "role", "guard"),
        "site_id": str(site.id) if site else None,
        "site_name": site.name if site else "—",
        "org_location_id": str(org_loc.id) if org_loc else None,
        "org_location_name": org_loc.name if org_loc else "—",
        "shift_name": (getattr(shift, "name", None) or getattr(shift, "shift_name", None) or "—") if shift else "—",
        "latitude": float(log.latitude) if log.latitude is not None else None,
        "longitude": float(log.longitude) if log.longitude is not None else None,
        "is_auto_checkout": bool(log.is_auto_checkout),
        "auto_checkout_reason": log.auto_checkout_reason or "Boundary breach: moved outside site boundary",
    }


def build_auto_checkout_summary(qs) -> dict:
    """KPI summary cards for dashboard and export headers."""
    total = qs.count()
    unique_guards = qs.values("guard_id").distinct().count()
    unique_sites = qs.values("site_id").exclude(site_id__isnull=True).distinct().count()
    return {
        "total_auto_checkouts": total,
        "unique_guards": unique_guards,
        "unique_sites": unique_sites,
    }


AUTO_CHECKOUT_EXPORT_HEADERS = [
    "#",
    "Date & Time",
    "Guard",
    "Emp Code",
    "Role",
    "Site",
    "Location",
    "Shift",
    "Breach GPS",
]


def resolve_auto_checkout_export_meta(request, requesting_user):
    from patrol_backend.utils.timezone_utils import (
        get_user_now,
        get_user_timezone_from_request,
        to_user_timezone,
    )

    if not request:
        user_tz = getattr(requesting_user, "timezone", None)
        return {
            "org_name": getattr(requesting_user, "location", None) and requesting_user.location.name or "—",
            "site_name": None,
            "range_start": None,
            "range_end": None,
            "printed_at": to_user_timezone(timezone.now(), requesting_user),
            "user_tz": user_tz,
        }

    start_date_str = request.query_params.get("start_date")
    end_date_str = request.query_params.get("end_date")
    location_id = request.query_params.get("location_id")
    site_id = request.query_params.get("site_id")

    user_tz = get_user_timezone_from_request(request, location_id=location_id)
    today = get_user_now(user_tz).date()

    range_start, range_end = None, None
    if start_date_str and end_date_str:
        try:
            range_start = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            range_end = datetime.strptime(end_date_str, "%Y-%m-%d").date()
        except ValueError:
            range_start, range_end = today, today
    elif start_date_str:
        try:
            range_start = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            range_end = range_start
        except ValueError:
            pass

    org_name = "—"
    if location_id and location_id not in ("All", "all", "null"):
        from scheduler.models import Location

        loc = Location.objects.filter(id=location_id, is_deleted=False).first()
        if loc:
            org_name = loc.name
    elif getattr(requesting_user, "location", None):
        org_name = requesting_user.location.name or "—"

    site_name = None
    if site_id and site_id not in ("All", "all", "null"):
        from scheduler.models import LocationSite

        site = LocationSite.objects.filter(id=site_id).select_related("location").first()
        if site:
            site_name = site.name
            if org_name == "—" and site.location:
                org_name = site.location.name

    return {
        "org_name": org_name,
        "site_name": site_name,
        "range_start": range_start,
        "range_end": range_end,
        "printed_at": to_user_timezone(timezone.now(), user_tz),
        "user_tz": user_tz,
    }


# =============================================================================
# EXCEL GENERATOR (openpyxl) — GTMS Standard Report Layout
# =============================================================================

def generate_auto_checkout_excel(queryset, requesting_user, request=None, context: Optional[dict] = None) -> io.BytesIO:
    """Produce formatted Excel workbook of auto-checkout events matching standard GTMS layout."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Auto-Checkout Report"
    ws.views.sheetView[0].showGridLines = True

    meta = resolve_auto_checkout_export_meta(request, requesting_user)
    user_tz = meta.get("user_tz") or requesting_user

    # Row 1: Header Row
    headers = AUTO_CHECKOUT_EXPORT_HEADERS
    ws.append(headers)

    header_font = Font(name="Calibri", size=10, bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 24
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    font_data = Font(name="Calibri", size=10)

    # Data Rows
    for idx, log in enumerate(queryset, start=1):
        local_time = to_user_timezone(log.timestamp, user_tz)
        time_str = local_time.strftime("%Y-%m-%d %H:%M:%S") if local_time else "—"

        guard = log.guard
        guard_name = getattr(guard, "name", None) or getattr(guard, "email", "Unknown")
        emp_code = getattr(guard, "employee_code", "") or "—"
        role = (getattr(guard, "role", "guard") or "guard").upper()
        site_name = log.site.name if log.site else "—"
        location_name = log.org_location.name if log.org_location else "—"
        shift_name = (getattr(log.shift, "name", None) or getattr(log.shift, "shift_name", None) or "—") if log.shift else "—"
        coords_str = (
            f"{float(log.latitude):.6f}, {float(log.longitude):.6f}"
            if (log.latitude is not None and log.longitude is not None)
            else "—"
        )

        row_data = [
            idx,
            time_str,
            guard_name,
            emp_code,
            role,
            site_name,
            location_name,
            shift_name,
            coords_str,
        ]
        ws.append(row_data)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = font_data
            cell.border = border_thin
            if col_idx in (1, 2, 4, 5, 9):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# =============================================================================
# PDF GENERATOR (ReportLab) — GTMS Standard Dashboard Report Layout
# =============================================================================

def generate_auto_checkout_pdf(queryset, requesting_user, request=None, context: Optional[dict] = None) -> io.BytesIO:
    """Produce formatted landscape PDF report of auto-checkout events matching GTMS standard report style."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    from patrol_backend.utils.pdf_report import draw_dashboard_pdf_header_and_footer

    meta = resolve_auto_checkout_export_meta(request, requesting_user)
    user_tz = meta.get("user_tz") or requesting_user

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=0.2 * inch,
        rightMargin=0.2 * inch,
        topMargin=1.2 * inch,
        bottomMargin=0.35 * inch,
    )
    doc.report_title = "Auto-Checkout Report"
    doc.org_name = meta["org_name"]
    doc.dept_label = meta.get("site_name") or meta["org_name"]
    doc.range_start = meta["range_start"]
    doc.range_end = meta["range_end"]
    doc.printed_at = meta["printed_at"]

    styles = getSampleStyleSheet()

    def _para_style(name, **kwargs):
        base = dict(
            leftIndent=0,
            rightIndent=0,
            firstLineIndent=0,
            spaceBefore=0,
            spaceAfter=0,
        )
        base.update(kwargs)
        return ParagraphStyle(name, **base)

    cell_style = _para_style(
        "AutoPdfCell",
        parent=styles["Normal"],
        fontSize=6.5,
        alignment=TA_LEFT,
        leading=8,
    )
    cell_center = _para_style(
        "AutoPdfCellCenter",
        parent=styles["Normal"],
        fontSize=6.5,
        alignment=TA_CENTER,
        leading=8,
    )
    header_style = _para_style(
        "AutoPdfHeader",
        parent=styles["Normal"],
        fontSize=6.5,
        fontName="Helvetica-Bold",
        alignment=TA_CENTER,
        leading=8,
    )

    data = [[Paragraph(h, header_style) for h in AUTO_CHECKOUT_EXPORT_HEADERS]]

    for idx, log in enumerate(queryset, start=1):
        local_time = to_user_timezone(log.timestamp, user_tz)
        time_str = local_time.strftime("%Y-%m-%d %H:%M") if local_time else "—"

        guard = log.guard
        guard_name = getattr(guard, "name", None) or getattr(guard, "email", "Unknown")
        emp_code = getattr(guard, "employee_code", "") or "—"
        role = (getattr(guard, "role", "guard") or "guard").upper()
        site_name = log.site.name if log.site else "—"
        location_name = log.org_location.name if log.org_location else "—"
        shift_name = (getattr(log.shift, "name", None) or getattr(log.shift, "shift_name", None) or "—") if log.shift else "—"
        coords_str = (
            f"{float(log.latitude):.4f}, {float(log.longitude):.4f}"
            if (log.latitude is not None and log.longitude is not None)
            else "—"
        )

        data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(time_str, cell_center),
            Paragraph(guard_name, cell_style),
            Paragraph(emp_code, cell_center),
            Paragraph(role, cell_center),
            Paragraph(site_name, cell_style),
            Paragraph(location_name, cell_style),
            Paragraph(shift_name, cell_style),
            Paragraph(coords_str, cell_center),
        ])

    if len(data) == 1:
        data.append(
            [Paragraph("No auto-checkout records found", cell_style)]
            + [""] * (len(AUTO_CHECKOUT_EXPORT_HEADERS) - 1)
        )

    # 9 columns: #, Date & Time, Guard, Emp Code, Role, Site, Location, Shift, Breach GPS
    ratios = [0.04, 0.13, 0.16, 0.09, 0.08, 0.15, 0.13, 0.10, 0.12]
    ratio_sum = sum(ratios)
    col_widths = [doc.width * (r / ratio_sum) for r in ratios]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )

    doc.build(
        [table],
        onFirstPage=draw_dashboard_pdf_header_and_footer,
        onLaterPages=draw_dashboard_pdf_header_and_footer,
    )
    buffer.seek(0)
    return buffer
